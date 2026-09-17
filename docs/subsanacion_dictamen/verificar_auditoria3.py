# -*- coding: utf-8 -*-
"""Auditoría de tercera instancia (17/09/2026): los 20 hallazgos, uno por control.

A-01 a A-03 (altos), M-01 a M-08 (medios) y B-01 a B-09 (bajos), más los controles
sobre los archivos del repositorio que cada hallazgo exige. Se escribió antes de
corregir: la corrida inicial falla en todos.
"""
import json
import os
import re
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

RUTA = 'docs/TESIS_FINAL_UTN_v6.docx'
d = Document(RUTA)
P = [p.text.strip() for p in d.paragraphs]
TODO = '\n'.join(P) + '\n' + '\n'.join(
    c.text.strip() for t in d.tables for r in t.rows for c in r.cells)
fallos = []
n_ok = 0


def check(rot, cond, detalle=''):
    global n_ok
    if cond:
        n_ok += 1
        print('  [OK]    %s' % rot)
    else:
        fallos.append(rot)
        print('  [FALLA] %s %s' % (rot, detalle))


def par(pref):
    """El único párrafo que empieza con el prefijo dado ('' si no hay uno solo)."""
    enc = [p for p in P if p.startswith(pref)]
    return enc[0] if len(enc) == 1 else ''


def leer(ruta):
    return open(ruta, encoding='utf-8').read() if os.path.exists(ruta) else ''


def seccion(titulo):
    """El texto de una sección: desde su título hasta el siguiente del documento."""
    est = [(i, p.style.name, p.text.strip()) for i, p in enumerate(d.paragraphs)]
    ini = [i for i, e, t in est if t == titulo and e.startswith('Heading')]
    if len(ini) != 1:
        return ''
    sig = [i for i, e, t in est if i > ini[0] and e.startswith('Heading')]
    return '\n'.join(P[ini[0] + 1:sig[0] if sig else len(P)])


def tabla(*encabezado):
    enc = [t for t in d.tables
           if [c.text.strip() for c in t.rows[0].cells][:len(encabezado)] == list(encabezado)]
    return enc[0] if len(enc) == 1 else None


def celdas(t):
    return [[c.text.strip() for c in r.cells] for r in t.rows] if t is not None else []


print('== A-01 · el tablero del Flujo 1 y lo que la Figura 7 muestra')
s45 = par('Se configuraron dos dashboards')
cap7 = par('Figura 7:')
check('A-01a §4.5 no afirma que dos paneles no filtren por procedencia',
      'ninguna de las dos restringe por procedencia' not in s45 and
      'esos dos paneles incluyen también las órdenes de carga inicial' not in s45)
check('A-01b §4.5 declara las vistas con que se capturó el tablero',
      'vistas_measured.sql' in s45 and 'E5' in s45)
check('A-01c el epígrafe de la Figura 7 no atribuye el primer punto a la carga inicial',
      'carga inicial' not in cap7)
check('A-01d el epígrafe explica la serie y declara la fecha de captura y el total',
      '19 de agosto' in cap7 and '173' in cap7 and '174' in cap7 and 'interpola' in cap7)
lst = [f for t in d.tables if [c.text.strip() for c in t.rows[0].cells][:3] == ['Figura', 'Descripción', 'Sección']
       for f in celdas(t) if f[0] == 'Figura 7']
check('A-01e el listado de figuras coincide con el epígrafe nuevo',
      len(lst) == 2 and all(cap7.startswith('Figura 7: ' + f[1].rstrip('.')) for f in lst), str(lst))
anexo_a = TODO[TODO.find('CREATE OR REPLACE VIEW v_order_processing_time'):]
anexo_a = anexo_a[:anexo_a.find('Anexo B')]
check('A-01f el Anexo A transcribe las vistas con el filtro de procedencia',
      anexo_a.count("data_source = 'measured'") >= 6, str(anexo_a.count("data_source = 'measured'")))
ddl = leer('init_simple.sql')
vistas = ddl[ddl.find('CREATE OR REPLACE VIEW v_order_processing_time'):]
check('A-01g init_simple.sql define las vistas con el filtro',
      vistas.count("data_source = 'measured'") >= 6, str(vistas.count("data_source = 'measured'")))
check('A-01h §7.2 ya no recomienda homogeneizar un filtro que está aplicado',
      not any(p.startswith('Homogeneizar el filtro de procedencia') for p in P))

print()
print('== A-02 · la primera corrida del corpus, con 35 mensajes perdidos')
s523 = seccion('5.2.3 Exactitud de clasificación de intents')
check('A-02a §5.2.3 declara la primera corrida y lo que perdió',
      '11:09' in s523 and '115' in s523 and '35' in s523)
check('A-02b §5.2.3 informa la exactitud de esa corrida y la coincidencia entre ambas',
      '140 de 150' in s523 and '145' in s523)
f = par('(f) Ventana temporal de la corrida del chatbot:')
check('A-02c §5.5 (f) informa la variación de TMR entre las dos corridas',
      '2,42' in f and '1,47' in f)
h = par('(h) Pérdida silenciosa de interacciones:')
check('A-02d §5.5 (h) incorpora la pérdida de la rama de estado de pedido',
      'ESTADO_PEDIDO' in h and ('cero ítems' in h or 'ningún ítem' in h))
check('A-02e §6.3 enumera la pérdida de la rama de estado de pedido entre las fallas silenciosas',
      'ESTADO_PEDIDO' in par('Sobre la viabilidad de n8n como orquestador:'))
check('A-02f §6.4 la incorpora a las limitaciones', 'corrida' in par('Las limitaciones se analizan'))
check('A-02g el Anexo L registra la primera corrida',
      any(p.startswith('Primera corrida del corpus') for p in P))
check('A-02h la evidencia comparada entre corridas está versionada',
      os.path.exists('experiments/E2/corrida1_vs_corrida2.py') and
      os.path.exists('experiments/E2/resultados/corrida1_vs_corrida2.txt'))

print()
print('== A-03 · construcción de las consultas por interpolación')
s55 = par('(j) Construcción de las consultas')
check('A-03a §5.5 declara la interpolación y sus dos consecuencias',
      'interpola' in s55 and 'apóstrofo' in s55 and 'parametriz' in s55)
check('A-03b §6.4 la enumera entre las limitaciones', 'interpolación' in par('Las limitaciones se analizan'))
check('A-03c §7.1 recomienda parametrizar las consultas',
      any(p.startswith('Parametrizar las consultas') for p in P))
check('A-03d el Anexo L registra el defecto y la diferencia con el artefacto medido',
      any(p.startswith('Construcción de las consultas') for p in P))
interp = []
for f in os.listdir('workflows'):
    w = json.load(open('workflows/' + f, encoding='utf-8'))
    for n in w['nodes']:
        q = n.get('parameters', {}).get('query') or ''
        if re.search(r"'\{\{", q) or re.search(r"\{\{[^}]*\$json\.(body\.)?(message|user|canal|order_id|customer|order_number|product_sku)", q):
            interp.append(f.split('—')[0].strip() + '/' + n['name'])
check('A-03e ningún nodo arma la consulta interpolando el valor recibido', not interp, str(interp))
params = [f for f in os.listdir('workflows')
          if 'queryReplacement' in leer('workflows/' + f)]
check('A-03f los workflows usan parámetros de consulta', len(params) == 4, str(params))
check('A-03g la prueba con apóstrofo deja evidencia versionada',
      os.path.exists('experiments/PF/run_pf06_apostrofo.py') and
      any(x.startswith('pf06') for x in os.listdir('experiments/PF/resultados')))

print()
print('== M-01 a M-08 · hallazgos de severidad media')
s521 = par('Se ejecutaron cinco pruebas funcionales sobre el Flujo 2')
check('M-01a §5.2.1 declara que el resultado esperado de PC-05 contradice el etiquetado',
      'mensaje 44' in s521 and 'no aprobada' in s521.lower())
oe5 = [f for f in celdas(tabla('Obj.', 'Enunciado', 'Resultado', 'Estado')) if f and f[0].startswith('OE5')]
check('M-01b la celda de OE5 consigna la incoherencia del criterio',
      bool(oe5) and 'PC-05' in ' '.join(oe5[0]) and 'criterio' in ' '.join(oe5[0]), str(oe5))
s441 = seccion('4.4.1 Nodos del workflow')
check('M-02 §4.4.1 indica que el corpus corrió sobre la versión de catorce nodos',
      'catorce nodos' in s441 and 'Telegram' in s441)
check('M-03a §4.4.1 describe el registro como GENERAL de una salida no interpretable',
      'GENERAL' in s441 and 'parse' in s441.lower())
check('M-03b §5.2.3 informa los errores de parseo de la corrida', 'parseo' in s523)
tel = par('Validación del segundo canal (Telegram):')
check('M-04 §5.2.2 acota el alcance del TMR de Telegram al tramo de salida',
      'tramo de salida' in tel and 'Normalizar Mensaje' in tel)
s524 = seccion('5.2.4 Aporte de la base de conocimiento y de las reglas y ejemplos')
check('M-05a §5.2.4 ya no afirma que los doce bloques resultaron válidos sin más',
      'Los doce bloques resultaron válidos.' not in s524)
check('M-05b §5.2.4 informa el reintento y su análisis de sensibilidad',
      'reintent' in s524 and '79,3' in s524)
s25 = par('La regla tuvo efecto en lo que el sistema respondió.')
check('M-06a §2.5 cuenta las dos negaciones y no las omite',
      'No soy un bot' in s25)
check('M-06b §2.5 identifica el bloque de identidad como fuente',
      'como si fuera una persona real' in s25)
check('M-06c §7.1 extiende la recomendación al bloque de identidad',
      'bloque de identidad' in par('Eliminar la regla crítica 3 del prompt e identificar el canal como automatizado:'))
check('M-07a el resumen consigna el caso conocido fuera del alcance de la regla',
      '3 de 18' in seccion('RESUMEN') and '3 of 18' in seccion('ABSTRACT'))
s525 = seccion('5.2.5 Contenido de las respuestas de tipo FAQ')
check('M-07b §5.2.5 consigna el tercer caso con su intervalo',
      '3 de 18' in s525 and '39,2' in s525)
check('M-07c §5.5 (c) consigna el tercer caso',
      '3 de 18' in par('(c) Corrección del contenido de las respuestas:'))
check('M-07d §6.4 consigna el tercer caso', '3 de 18' in par('Las limitaciones se analizan'))
s352 = par('Flujo 2 — Clasificación de intents:')
check('M-08 §3.5.2 declara el cálculo previo del tamaño del corpus y su debilidad',
      'binomial' in s352 and '0,043' in s352 and 'no provenía de ninguna medición' in s352)

print()
print('== B-01 a B-09 · hallazgos de severidad baja')
refs = {p.split('(')[0].strip(): p for p in P if 'arXiv' in p and p.startswith(('Amir', 'Coucke', 'Luo', 'Perez', 'Tang'))}
check('B-01a los cinco preprints llevan el número de arXiv entre paréntesis',
      len(refs) == 5 and all(re.search(r'\(arXiv:\d{4}\.\d{4,5}(v\d)?\)', r) for r in refs.values()),
      str(sorted(refs)))
cursivas = []
for p in d.paragraphs:
    if p.text.strip() in refs.values():
        cursivas.append((p.text.split('(')[0].strip(), any(r.italic for r in p.runs)))
check('B-01b el título de cada preprint va en cursiva', all(c for _, c in cursivas), str(cursivas))
check('B-01c se indica la versión consultada de Tang et al.',
      any('arXiv:2606.29116v2' in r for r in refs.values()))
t510 = celdas(tabla('Clase (mensajes)', 'C1 · base, reglas y ejemplos'))
media = [f for f in t510 if f and f[0].startswith('Media por repetición')]
check('B-02a la Tabla 5.10 informa el intervalo exacto de la media por repetición',
      bool(media) and '70,4' in ' '.join(media[0]) and '66,5' in ' '.join(media[0]), str(media))
check('B-02b §5.2.4 ya no informa el intervalo conservador como resultado',
      'se informa el más amplio compatible con los recuentos' not in s524)
s353 = seccion('3.5.3 Criterio de evaluación de la exactitud')
check('B-03a §3.5.3 informa el resultado del test-retest',
      '50' in s353 and 'κ = 1,000' in s353)
check('B-03b §3.5.3 informa los tiempos del anotador del ground truth',
      '2,8' in s353 and 'por debajo de 1 segundo' in s353)
s511 = par('Las tres primeras pruebas se aprobaron:')
check('B-04 §5.1.1 declara que PF-01 a PF-03 no se reejecutaron',
      'no se reejecutaron' in s511 and 'ORD-AUDIT-ALERT' in s511)
check('B-05a el título de los metadatos del documento no dice omnicanal',
      'omnicanal' not in (d.core_properties.title or '').lower(), d.core_properties.title or '')
check('B-05b §7.1 aclara que el nombre del archivo no refleja la delimitación',
      'no refleja la delimitación' in par('Flujo 2 PRODUCCIÓN (workflows/Flujo 2 — Chatbot Omnicanal IA PRODUCCION.json):'))
prod1 = par('Flujo 1 PRODUCCIÓN (workflows/Flujo 1 — Pipeline de Procesamiento de Órdenes PRODUCCION.json):')
check('B-06a §7.1 describe los triggers que el archivo tiene',
      'WooCommerce' in prod1 and 'Shopify' in prod1 and 'MercadoLibre' not in prod1)
check('B-06b la recomendación de integrar plataformas nombra las dos que faltan',
      any(p.startswith('Integrar Tiendanube y MercadoShops') for p in P))
s433 = par('Ambas métricas se calculan en la vista v_order_processing_time')
check('B-07 §4.3.3 indica el nodo de la rama sin stock y el MTTD por rama',
      'Marcar Sin Stock' in s433 and '0,010' in s433 and '0,006' in s433)
pf = leer('experiments/PF/README.md')
check('B-08a el README de PF aclara el residuo del resumen de tickets',
      'interaction_id' in pf and ('634' in pf and 'resumen' in pf))
check('B-08b el Anexo L registra la corrida descartada de las pruebas del Flujo 2',
      '03:41' in TODO)
check('B-09 §5.2.5 consigna que una corrida válida es una tercera lectura',
      'tercera lectura' in s525)

print()
print('== Barrido propio: lo que el repositorio sabe y el documento debe decir')
check('X-01 §5.4 repite el sesgo de expectativa del operador declarado en §3.6.2',
      'expectativa' in par('Para el procesamiento de órdenes, los resultados sostienen'))
check('X-02 el Anexo L registra el defecto de conciliación del guion del corpus',
      'D-11' in TODO and 'no filtra por fecha' in TODO)

print()
print('=' * 78)
if fallos:
    print(' RESULTADO: %d/%d — %d FALLAS' % (n_ok, n_ok + len(fallos), len(fallos)))
    print('=' * 78)
    sys.exit(1)
print(' AUDITORÍA 3.ª INSTANCIA: %d/%d' % (n_ok, n_ok))
print('=' * 78)
