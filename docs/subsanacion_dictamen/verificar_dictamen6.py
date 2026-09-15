# -*- coding: utf-8 -*-
"""Dictamen del 14/09, segunda pasada: reestructuración, estado del arte y recomendadas pendientes.

Complementa a verificar_dictamen5.py. Además de un control por pedido, incluye
uno general que la reestructuración vuelve imprescindible: toda remisión a una
sección, capítulo o tabla apunta a una que existe. Se escribió antes de corregir.
"""
import sys
import io
import re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

d = Document('docs/TESIS_FINAL_UTN_v6.docx')
P = [p.text.strip() for p in d.paragraphs]
HEAD = [p.text.strip() for p in d.paragraphs if p.style.name.startswith('Heading') and p.text.strip()]
TXT_P = '\n'.join(P)
CELDAS = '\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
TODO = TXT_P + '\n' + CELDAS
i_ref = next(k for k, t in enumerate(P) if t.startswith('CAPÍTULO 8'))
i_anx = next(k for k, t in enumerate(P) if t.startswith('CAPÍTULO 9'))
CAP1_7 = '\n'.join(P[:i_ref])
REFS = [t for t in P[i_ref + 1:i_anx] if t]

fallos = []
n_ok = 0


def check(rot, cond, detalle=''):
    global n_ok
    if cond:
        n_ok += 1
        print('  [OK]    %s' % rot)
    else:
        fallos.append(rot + ((' — ' + detalle) if detalle else ''))
        print('  [FALLA] %s %s' % (rot, detalle))


def bloque(desde, hasta):
    try:
        i = next(k for k, t in enumerate(P) if t.startswith(desde))
        j = next(k for k, t in enumerate(P) if k > i and t.startswith(hasta))
    except StopIteration:
        return ''
    return '\n'.join(P[i:j])


def seccion(titulo):
    print()
    print('=' * 78)
    print(' ' + titulo)
    print('=' * 78)


# ============================================================================
seccion('A. REMISIONES: TODA SECCIÓN, CAPÍTULO Y TABLA CITADA EXISTE')
numeros = set()
for h in HEAD:
    m = re.match(r'^(\d+(?:\.\d+)+)\s', h)
    if m:
        numeros.add(m.group(1))
caps = {m.group(1) for h in HEAD for m in [re.match(r'^CAPÍTULO (\d+)', h)] if m}
epigrafes = {m.group(1) for t in P for m in [re.match(r'^Tabla ([\dA-K]+\.\d+):', t)] if m}
rotas_sec, rotas_tab, rotas_cap = set(), set(), set()
for m in re.finditer(r'(?:Secci[oó]n(?:es)?|§)\s*(\d+\.\d+(?:\.\d+)?(?:(?:,\s*|\s+y\s+|\s+a\s+)\d+\.\d+(?:\.\d+)?)*)', CAP1_7 + '\n' + CELDAS):
    for n in re.findall(r'\d+\.\d+(?:\.\d+)?', m.group(1)):
        if n not in numeros:
            rotas_sec.add(n)
for m in re.finditer(r'Tablas?\s+([\dA-K]+\.\d+(?:(?:,\s*|\s+y\s+|\s+a\s+)[\dA-K]+\.\d+)*)', CAP1_7):
    for n in re.findall(r'[\dA-K]+\.\d+', m.group(1)):
        if n not in epigrafes:
            rotas_tab.add(n)
for m in re.finditer(r'Cap[ií]tulo (\d+)', CAP1_7):
    if m.group(1) not in caps:
        rotas_cap.add(m.group(1))
check('A1  toda «Sección» citada existe', not rotas_sec, str(sorted(rotas_sec)))
check('A2  toda «Tabla» citada existe', not rotas_tab, str(sorted(rotas_tab)))
check('A3  todo «Capítulo» citado existe', not rotas_cap, str(sorted(rotas_cap)))
nums2 = [h.split()[0] for h in HEAD if re.match(r'^2\.\d', h)]
check('A4  la numeración del Capítulo 2 es correlativa',
      nums2 == ['2.1', '2.1.1', '2.1.2', '2.1.3', '2.2', '2.2.1', '2.2.2', '2.2.3', '2.2.4', '2.2.5', '2.3', '2.3.1',
                '2.3.2', '2.4', '2.4.1', '2.4.2', '2.4.3', '2.4.4', '2.4.5', '2.4.6', '2.5'], str(nums2))
nums3 = [h.split()[0] for h in HEAD if re.match(r'^3\.\d', h)]
check('A5  la numeración del Capítulo 3 es correlativa',
      nums3 == ['3.1', '3.2', '3.3', '3.4', '3.4.1', '3.4.2', '3.4.3', '3.5', '3.5.1', '3.5.2', '3.5.3', '3.5.4', '3.5.5',
                '3.5.6', '3.5.7', '3.5.8', '3.6', '3.6.1', '3.6.2', '3.6.3', '3.6.4'], str(nums3))
nums4 = [h.split()[0] for h in HEAD if re.match(r'^4\.6', h)]
check('A6  §4.6 sin subsecciones de diseño', nums4 == ['4.6'], str(nums4))
lst = [[c.text.strip() for c in r.cells] for t in d.tables for r in t.rows
       if r.cells[0].text.strip().startswith('Tabla ') and len(t.columns) == 3 and t.rows[0].cells[0].text.strip() == 'Tabla']
orden = [f[0] for f in lst]
cap_de = {m.group(1): m.group(2) for t in P for m in [re.match(r'^Tabla ([\dA-K]+\.\d+): (.*?)\.?$', t)] if m}
check('A7  el listado de tablas sigue el orden del documento y sus epígrafes',
      orden[:6] == ['Tabla 2.1', 'Tabla 3.1', 'Tabla 3.2', 'Tabla 3.3', 'Tabla 3.4', 'Tabla 3.5']
      and all(cap_de.get(f[0][6:], '').startswith(f[1]) for f in lst), str(orden[:8]))   # algunos epígrafes agregan detalle
check('A8  el listado ubica cada tabla en su sección',
      dict((f[0], f[2]) for f in lst).get('Tabla 2.1') == '2.4.6'
      and dict((f[0], f[2]) for f in lst).get('Tabla 3.4') == '3.5.8')

# ============================================================================
seccion('B. REESTRUCTURACIÓN')
check('B1  la infraestructura pasó del Capítulo 2 al 3',
      all(h in HEAD for h in ('3.4.1 Contenedores Docker y reproducibilidad', '3.4.2 PostgreSQL como base de datos de métricas',
                              '3.4.3 Grafana para observabilidad operativa'))
      and not any(h.startswith('2.4 Infraestructura') for h in HEAD)
      and 'Corresponde declarar de antemano la función de esta sección' not in TODO)
b358 = bloque('3.5.8', '3.6 Amenazas')
check('B2  el diseño de las pruebas pasó a §3.5.8',
      '3.5.8 Diseño de las pruebas funcionales, de carga y de concurrencia' in HEAD
      and 'Tabla 3.4: Pruebas funcionales del Flujo 1.' in b358 and 'Tabla 3.5: Pruebas funcionales del Flujo 2.' in b358
      and 'El segundo ensayo (E1.b)' in b358 and 'Myers et al., 2011' in b358)
b46 = bloque('4.6 Entorno', 'CAPÍTULO 5')
check('B3  §4.6 queda como entorno y evidencia', '4.6 Entorno y evidencia de las pruebas' in HEAD
      and 'Intel Core i5' in b46 and 'Figura 8:' in b46 and 'no son apareadas' not in b46)
b522 = bloque('5.2.2', '5.2.3')
check('B4  la no comparabilidad entre canales se discute en §5.2.2', 'no son apareadas' in b522 and 'prompt reducido' in b522)
b36 = bloque('3.6 Amenazas', 'CAPÍTULO 4')
b443 = bloque('4.4.3', '4.4.4')
check('B5  la narración de la auditoría pasó a amenazas a la validez',
      'La auditoría fue correcta en su método' in b36 and 'La auditoría fue correcta en su método' not in b443
      and 'trazabilidad_versiones.txt' in b443)
b71 = bloque('7.1 ', '7.2 ')
b72 = bloque('7.2 ', 'CAPÍTULO 8')
check('B6  §7.1 contiene solo recomendaciones de producción',
      not any(x in b71 for x in ('Capturar la marca de recepción', 'Homogeneizar el filtro', 'Aislar el componente de red'))
      and all(x in b72 for x in ('Capturar la marca de recepción', 'Homogeneizar el filtro', 'Aislar el componente de red'))
      and '7.2 Recomendaciones para la investigación y líneas futuras' in HEAD)
check('B7  §7.1 incorpora Tiendanube, protección de datos y alineación de reglas',
      all(x in b71 for x in ('Tiendanube', 'protección de datos personales', 'Alinear las reglas del prompt con el contenido de la base')))
defensivo = re.findall(r'[Cc]orresponde (?:declarar|precisar|señalar|consignar|subrayar|aclarar|revisar|explicitar|una precisión|una salvedad|por lo tanto)', CAP1_7)
conviene = re.findall(r'\b[Cc]onviene\b', CAP1_7)
sedeclara = re.findall(r'\bse declara\b', CAP1_7)
check('B8  metadiscurso defensivo reducido', not defensivo and len(conviene) <= 2 and len(sedeclara) <= 3,
      'corresponde+verbo %d, conviene %d, se declara %d' % (len(defensivo), len(conviene), len(sedeclara)))
b61 = bloque('6.1', '6.2')
check('B9  §6.1 no repite las salvedades del factor', 'operan en direcciones opuestas' not in b61)

# ============================================================================
seccion('C. ESTADO DEL ARTE Y MARCO TEÓRICO')
b24 = bloque('2.4 Estado del arte', '2.5 Tratamiento')
check('C1  la búsqueda se actualizó con los términos nuevos', 'retrieval-augmented generation' in bloque('2.4 Estado del arte', '2.4.1'))
check('C2  subsección de anclaje y alucinación', '2.4.3 Anclaje en bases de conocimiento y alucinación' in HEAD)
check('C3  la contribución se presenta como integración',
      '2.4.6 Contribución de este trabajo' in HEAD and 'La contribución es de integración' in bloque('2.4.6', '2.5 Tratamiento'))
nuevas = [('Lewis et al. (2020)', 'Lewis, P., Perez, E.'), ('Fan et al. (2024)', 'Fan, W., Ding, Y.'),
          ('Magesh et al. (2025)', 'Magesh, V.'), ('Larsen et al. (2026)', 'Larsen, A. G.'),
          ('Moffatt v. Air Canada', 'Moffatt v. Air Canada'), ('Arora et al. (2024)', 'Arora, G.'),
          ('Liu et al. (2024)', 'Liu, N. F.'), ('Sclar et al. (2024)', 'Sclar, M.'), ('Ouyang et al. (2025)', 'Ouyang, S.'),
          ('Gatt y Krahmer (2018)', 'Gatt, A., & Krahmer, E.'), ('Artstein y Poesio (2008)', 'Artstein, R., & Poesio, M.'),
          ('Bock y Frank (2021)', 'Bock, A. C., & Frank, U.'), ('Sahay et al. (2020)', 'Sahay, A.'),
          ('Rokis y Kirikova (2022)', 'Rokis, K., & Kirikova, M.')]
for cita, ref in nuevas:
    parentetica = re.sub(r' \((\d{4})\)$', r', \1', cita)   # «Autor (año)» o «(Autor, año)»
    check('C4  %-26s citado y en la lista' % cita, (cita in CAP1_7 or parentetica in CAP1_7) and any(r.startswith(ref) for r in REFS))
check('C5  Huang et al. en su versión publicada', 'Huang et al., 2023' not in TODO and 'Huang et al. (2023)' not in TODO
      and any(r.startswith('Huang, L.') and '10.1145/3703155' in r for r in REFS))
recientes = [c for c, _ in nuevas if re.search(r'202[4-6]', c) or c.startswith('Moffatt')]
check('C6  al menos seis fuentes de 2024 a 2026 sobre LLM, anclaje o alucinación en §2.2 y §2.4',
      sum(1 for c in recientes if c in bloque('2.2 Inteligencia', '2.3 Comunicación') + b24) >= 6, str(recientes))
t21 = [[c.text.strip() for c in r.cells] for t in d.tables for r in t.rows if t.rows[0].cells[0].text.strip() == 'Antecedente']
check('C7  la Tabla 2.1 incorpora los antecedentes recientes',
      all(any(f[0].startswith(x) for f in t21) for x in ('Arora et al. (2024)', 'Magesh et al. (2025)', 'Larsen et al. (2026)')))
check('C8  §2.2.4 distingue inyección de contexto y RAG',
      '2.2.4 Inyección de contexto, generación aumentada por recuperación y alucinación' in HEAD
      and 'recuperación de contexto' not in TODO and 'inyección de contexto' in b443)
check('C9  §2.2.5 funda evaluación de texto, no determinismo y acuerdo',
      '2.2.5 Evaluación de texto generado, no determinismo y acuerdo entre evaluadores' in HEAD)
check('C10 §2.1.2 incorpora literatura académica low-code', 'Bock y Frank (2021)' in bloque('2.1.2', '2.1.3'))
check('C11 §2.3.2 no presenta el TMR como efectividad', '2.3.2 TMR y capacidad de respuesta' in HEAD
      and 'métrica operativa central para evaluar la efectividad' not in TODO)
check('C12 «tiempo marginal» en lugar de «costo marginal»', not re.search(r'costos? marginal', TODO, re.I)
      and 'tiempo marginal de procesamiento' in bloque('3.5.5', '3.5.6'))
check('C13 §5.2.5 conecta el hallazgo con la literatura', 'Magesh et al., 2025' in bloque('5.2.5', '5.3 '))

# ============================================================================
seccion('D. RECOMENDADAS PENDIENTES')
check('D1  §3.5.2 declara que la suficiencia muestral es posterior al resultado', 'posterior al resultado' in bloque('3.5.2', '3.5.3'))
check('D2  §3.6.3 corrige la mitigación de validez externa',
      'permite replicar el estudio con datos reales como trabajo futuro' not in TODO
      and 'no convierte los resultados en generalizables' in b36)
CAP7 = bloque('CAPÍTULO 7', 'CAPÍTULO 8')
check('D3  el instrumento de contenido se pilotea y se extiende a GENERAL y ESTADO_PEDIDO',
      'pilotearse' in CAP7 and 'GENERAL y ESTADO_PEDIDO' in CAP7)
check('D4  la línea de ajuste fino considera la ambigüedad de las etiquetas', 'ambigüedad de las etiquetas' in CAP7)

print()
print('=' * 78)
if fallos:
    print(' RESULTADO: %d/%d — %d FALLAS' % (n_ok, n_ok + len(fallos), len(fallos)))
    print('=' * 78)
    sys.exit(1)
print(' DICTAMEN DEL 14/09, SEGUNDA PASADA: %d/%d' % (n_ok, n_ok))
print('=' * 78)
