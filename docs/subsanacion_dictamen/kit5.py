# -*- coding: utf-8 -*-
"""Utilidades para la pasada del dictamen del 14/09 (fix_d5_*.py).

Todo se localiza por el texto con que empieza el párrafo y cada operación
exige encontrar exactamente lo que espera: si el documento no está en el
estado previsto, el guion aborta antes de guardar.
"""
import os
import sys
import copy

sys.path.insert(0, 'docs/subsanacion_dictamen')
from docxkit import replace_everywhere, set_cell  # noqa: E402,F401
from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

RUTA = 'docs/TESIS_FINAL_UTN_v6.docx'


def abrir():
    if any(f.startswith('~$') for f in os.listdir('docs')):
        sys.exit('ERROR: Word tiene abierto un documento en docs/. Cerralo primero.')
    return Document(RUTA)


class Kit:
    def __init__(self, d):
        self.d = d
        self.hechos = 0

    # ---------------------------------------------------------- localizar
    def pars(self, pref):
        return [p for p in self.d.paragraphs if p.text.strip().startswith(pref)]

    def par(self, pref):
        enc = self.pars(pref)
        assert len(enc) == 1, 'se esperaba 1 párrafo que empiece con %r, hay %d' % (pref[:70], len(enc))
        return enc[0]

    def tabla(self, encabezado):
        """La única tabla cuya primera fila empieza con las celdas dadas."""
        enc = [t for t in self.d.tables
               if [c.text.strip() for c in t.rows[0].cells][:len(encabezado)] == list(encabezado)]
        assert len(enc) == 1, 'se esperaba 1 tabla con encabezado %r, hay %d' % (encabezado, len(enc))
        return enc[0]

    def fila(self, tabla, primera_celda):
        enc = [i for i, r in enumerate(tabla.rows) if r.cells[0].text.strip() == primera_celda]
        assert len(enc) == 1, 'se esperaba 1 fila %r, hay %d' % (primera_celda, len(enc))
        return enc[0]

    # ---------------------------------------------------------- escribir
    def reescribir(self, pref, texto, etiqueta=None):
        """Reemplaza el texto del párrafo. Con `etiqueta`, la escribe en el
        primer run (que conserva su formato, p. ej. negrita) y el resto en un
        run sin negrita."""
        p = pref if isinstance(pref, Paragraph) else self.par(pref)
        assert 'w:drawing' not in p._p.xml, 'el párrafo lleva imagen'
        runs = p.runs
        assert runs, 'párrafo sin runs'
        if etiqueta is None:
            runs[0].text = texto
            for r in runs[1:]:
                r.text = ''
        else:
            assert texto.startswith(etiqueta), 'el texto no empieza con la etiqueta'
            runs[0].text = etiqueta
            if len(runs) > 1:
                runs[1].text = texto[len(etiqueta):]
                runs[1].bold = None
                for r in runs[2:]:
                    r.text = ''
            else:
                nuevo = p.add_run(texto[len(etiqueta):])
                if runs[0]._r.rPr is not None:
                    nuevo._r.insert(0, copy.deepcopy(runs[0]._r.rPr))
                nuevo.bold = None
        self.hechos += 1
        return p

    def reemplazo(self, old, new, n=1):
        k = replace_everywhere(self.d, old, new)
        assert k == n, 'se esperaban %d reemplazos de %r, hubo %d' % (n, old[:70], k)
        self.hechos += 1

    def celda(self, tabla, fila, col, texto):
        set_cell(tabla, fila, col, texto)
        self.hechos += 1

    def insertar_despues(self, ancla, modelo, texto, etiqueta=None):
        """Párrafo nuevo con el formato de `modelo`, inmediatamente después de
        `ancla`. No copia marcadores ni otros hijos del modelo."""
        ancla = ancla if isinstance(ancla, Paragraph) else self.par(ancla)
        modelo = modelo if isinstance(modelo, Paragraph) else self.par(modelo)
        el = OxmlElement('w:p')
        if modelo._p.pPr is not None:
            pPr = copy.deepcopy(modelo._p.pPr)
            for s in pPr.findall(qn('w:sectPr')):
                pPr.remove(s)
            el.append(pPr)
        ancla._p.addnext(el)
        p = Paragraph(el, ancla._parent)
        rm = modelo.runs
        partes = [(texto, rm[0] if rm else None)] if etiqueta is None else \
            [(etiqueta, rm[0] if rm else None), (texto[len(etiqueta):], rm[1] if len(rm) > 1 else None)]
        for i, (t, run_modelo) in enumerate(partes):
            r = p.add_run(t)
            if run_modelo is not None and run_modelo._r.rPr is not None:
                r._r.insert(0, copy.deepcopy(run_modelo._r.rPr))
            if etiqueta is not None and i == 1:
                r.bold = None
        self.hechos += 1
        return p

    def eliminar(self, pref):
        p = pref if isinstance(pref, Paragraph) else self.par(pref)
        p._p.getparent().remove(p._p)
        self.hechos += 1

    def guardar(self):
        self.d.save(RUTA)
        print('guardado: %d operaciones' % self.hechos)
