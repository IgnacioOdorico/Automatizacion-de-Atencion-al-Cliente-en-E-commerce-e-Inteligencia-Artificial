# -*- coding: utf-8 -*-
"""Residuos de la auditoría de 2.ª instancia encontrados al releerla completa.

  1. B-08 (segunda parte, abierta): el canvas de la Figura 6 rotula el nodo 17
     «Enviar Respuesta (Producción)» y la Tabla 4.7 lo llamaba «Enviar Respuesta
     (canal simulado)». La tabla adopta el rótulo del canvas y aclara que no
     designa los workflows de producción del Capítulo 7.
  2. Introducido por la corrección de E7: §4.4.3 decía «workflow de diecinueve
     nodos». El archivo tiene 19 elementos, pero uno es una nota adhesiva; son
     18 nodos, los que cuentan la Tabla 4.7 y la auditoría (M-09).
  3. Introducido por E7: la Figura 5 rotulaba el modelo «gpt-4o-mini ·
     zero-shot», que solo vale para la condición de ablación. Se corrige el SVG,
     se regenera el PNG con cairosvg (que reproduce el PNG anterior al píxel) y
     se reemplaza la imagen dentro del .docx, con las mismas dimensiones.

NO es idempotente. Se ejecuta desde la raíz del repositorio.
"""
import sys
import os
import re
import io

sys.path.insert(0, 'docs/subsanacion_dictamen')
from docxkit import replace_everywhere, set_cell  # noqa: E402
from docx import Document  # noqa: E402
import cairosvg  # noqa: E402
from PIL import Image  # noqa: E402

if any(f.startswith('~$') for f in os.listdir('docs')):
    sys.exit('ERROR: Word tiene abierto un documento en docs/.')

RUTA = 'docs/TESIS_FINAL_UTN_v6.docx'
SVG = 'docs/figuras_v6/diagrama_flujo2.svg'
PNG = 'docs/figuras_v6/diagrama_flujo2.png'
d = Document(RUTA)

# ------------------------------------------------------------------ 1. B-08
filas = [(t, i) for t in d.tables for i, r in enumerate(t.rows)
         if r.cells[0].text.strip() == '17' and r.cells[1].text.strip() == 'Enviar Respuesta (canal simulado)']
assert len(filas) == 1, 'no se encontró la fila 17 de la Tabla 4.7 en su estado anterior'
t, i = filas[0]
set_cell(t, i, 1, 'Enviar Respuesta (Producción)')
set_cell(t, i, 3, 'Envía la respuesta por email: canal simulado en formato WhatsApp Cloud API con '
                  'entrega SMTP. El rótulo «(Producción)» es el que el nodo conserva en el canvas '
                  '(Figura 6); pertenece a este workflow de desarrollo, no a los workflows de '
                  'producción del Capítulo 7.')
print('1. Tabla 4.7, nodo 17: rótulo del canvas y aclaración')

# ------------------------------------------------------------------ 2. nodos
n = replace_everywhere(d, 'workflow de diecinueve nodos incorporado el 27 de agosto',
                          'workflow de dieciocho nodos incorporado el 27 de agosto')
assert n == 1, 'reemplazos: %d' % n
print('2. §4.4.3: dieciocho nodos')

# ------------------------------------------------------------------ 3. Figura 5
svg = io.open(SVG, encoding='utf-8').read()
viejo = '<text class="ns" x="910" y="453">gpt-4o-mini · zero-shot</text>'
nuevo = '<text class="ns" x="910" y="453">gpt-4o-mini · prompts en el Anexo H</text>'
assert svg.count(viejo) == 1
svg = svg.replace(viejo, nuevo)
io.open(SVG, 'w', encoding='utf-8', newline='\n').write(svg)
png = cairosvg.svg2png(bytestring=svg.encode('utf-8'))
assert Image.open(io.BytesIO(png)).size == Image.open(PNG).size, 'cambió el tamaño de la figura'
open(PNG, 'wb').write(png)

parte = None
for k, p in enumerate(d.paragraphs):
    if p.text.strip().startswith('Figura 5:'):
        rid = re.search(r'r:embed="(rId\d+)"', d.paragraphs[k - 1]._element.xml).group(1)
        parte = d.part.rels[rid].target_part
assert parte is not None
parte._blob = png
print('3. Figura 5: SVG corregido, PNG regenerado y reemplazado en el docx (%s)' % parte.partname)

d.save(RUTA)
print('guardado.')
