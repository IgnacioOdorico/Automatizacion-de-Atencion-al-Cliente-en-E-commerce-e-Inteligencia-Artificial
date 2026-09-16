# -*- coding: utf-8 -*-
"""Dictamen del 15/09: grupo A (antes del depósito) e inconsistencias remanentes.

Un control por cada pedido del grupo A y por cada fila de la tabla de
inconsistencias remanentes. Incluye además la corrección del origen del corpus,
que el dictamen no advirtió porque no consultó el repositorio: la tesis decía
que el equipo había escrito los 150 mensajes, y experiments/E2/README.md declara
que se redactaron con un asistente basado en un modelo de lenguaje.
Se escribió antes de corregir.
"""
import sys
import io
import re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

d = Document('docs/TESIS_FINAL_UTN_v6.docx')
P = [p.text.strip() for p in d.paragraphs]
TXT_P = '\n'.join(P)
CELDAS = '\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
TODO = TXT_P + '\n' + CELDAS
i_ref = next(k for k, t in enumerate(P) if t.startswith('CAPÍTULO 8'))
i_anx = next(k for k, t in enumerate(P) if t.startswith('CAPÍTULO 9'))
CAP1_7 = '\n'.join(P[:i_ref])
REFS = [t for t in P[i_ref + 1:i_anx] if t]

fallos = []
n_ok = 0


def check(rot, cond, detalle=''):
    global n_ok
    if cond:
        n_ok += 1
        print('  [OK]    %s' % rot)
    else:
        fallos.append(rot + ((' — ' + detalle) if detalle else ''))
        print('  [FALLA] %s %s' % (rot, detalle))


def bloque(desde, hasta):
    try:
        i = next(k for k, t in enumerate(P) if t.startswith(desde))
        j = next(k for k, t in enumerate(P) if k > i and t.startswith(hasta))
    except StopIteration:
        return ''
    return '\n'.join(P[i:j])


def tabla_con_epigrafe(prefijo):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefijo):
            sig = p._p.getnext()
            while sig is not None and sig.tag.endswith('}p') and not ''.join(sig.itertext()).strip():
                sig = sig.getnext()
            for t in d.tables:
                if t._tbl is sig:
                    return t
    return None


def filas(t):
    return [[c.text.strip() for c in r.cells] for r in t.rows] if t is not None else []


def seccion(titulo):
    print()
    print('=' * 78)
    print(' ' + titulo)
    print('=' * 78)


b25 = bloque('2.5 Tratamiento', 'CAPÍTULO 3')
b353 = bloque('3.5.3', '3.5.4')
b356 = bloque('3.5.6', '3.5.7')
b36 = bloque('3.6 Amenazas', 'CAPÍTULO 4')
b511 = bloque('5.1.1', '5.1.2')
b514 = bloque('5.1.4', '5.2 Resultados')
b524 = bloque('5.2.4', '5.2.5')
b525 = bloque('5.2.5', '5.3 ')
b53 = bloque('5.3 ', '5.4 ')
b54 = bloque('5.4 Análisis', '5.4.1')
b61 = bloque('6.1', '6.2')
b62 = bloque('6.2', '6.3')
b63 = bloque('6.3', '6.4')
b64 = bloque('6.4', 'CAPÍTULO 7')
b71 = bloque('7.1', '7.2')
anxH = bloque('Anexo H', 'Anexo I')
i_res = P.index('RESUMEN')
RES = ' '.join(P[i_res + 1:next(k for k, t in enumerate(P) if k > i_res and t.startswith('Palabras clave'))])
i_abs = P.index('ABSTRACT')
ABS = ' '.join(P[i_abs + 1:next(k for k, t in enumerate(P) if k > i_abs and t.startswith('Keywords'))])

# ============================================================================
seccion('A1. REGLA CRÍTICA 3: OCULTAR QUE EL ASISTENTE ES UNA IA')
check('A1a §2.5 ya no la describe como decisión de tono', 'no presentarse como un chatbot genérico' not in TODO)
check('A1b §2.5 transcribe la regla', 'NUNCA reveles que sos una IA' in b25)
check('A1c §2.5 la analiza frente al deber de informar, con fuente',
      'carácter automatizado' in b25 and 'Reglamento (UE) 2024/1689' in b25 and 'Moffatt' in b25)
check('A1d §2.5 declara la contradicción con la entrada «Soporte» de la base',
      'Nuestro sistema de IA te asiste' in b25)
i_regl = next((k for k, r in enumerate(REFS) if r.startswith('Reglamento (UE) 2024/1689')), -1)
check('A1e §2.5 informa qué respondió el sistema cuando se le preguntó si era una persona',
      'Soy una persona real' in b25 and 'identidad_asistente' in b25)
check('A1f la referencia del Reglamento está en la lista, entre Ramos De Santis y Rokis',
      i_regl > 0 and REFS[i_regl - 1].startswith('Ramos De Santis') and REFS[i_regl + 1].startswith('Rokis'))
check('A1g §4.4.3 la nombra entre las reglas críticas', 'no revelar que el asistente es una IA' in bloque('4.4.3', '4.4.4'))
check('A1h §7.1 recomienda eliminarla e identificar el canal como automatizado',
      'Eliminar la regla crítica 3' in b71 and 'automatizado' in b71)
check('A1i el Anexo H advierte la regla al transcribirla', 'regla crítica 3' in anxH)

# ============================================================================
seccion('A2. EJEMPLOS DEL PROMPT FRENTE A LA BASE')
check('A2a §5.2.5 declara el ejemplo de Córdoba con las dos cifras',
      'entre 3 y 5 días hábiles' in b525 and 'entre 3 y 7' in b525)
check('A2b §5.2.5 verifica los siete ejemplos con el mismo guion', 'siete ejemplos' in b525
      and 'verificar_ejemplos_prompt.py' in b525)
check('A2c §5.2.5 declara los compromisos sin cifra', 'hoy mismo' in b525 and 'de inmediato' in b525)
check('A2d §5.2.5 informa que el dato no llegó a las respuestas medidas', 'Tabla K.1' in b525
      and 'no figura en ninguna de las 45 respuestas' in b525)
check('A2e §7.1 suma los ejemplos a la recomendación de alineación',
      'Alinear las reglas del prompt con el contenido de la base' in b71 and 'Córdoba' in b71)
check('A2f el Anexo H advierte el ejemplo al transcribirlo', 'Córdoba' in anxH and 'Sección 5.2.5' in anxH)

# ============================================================================
seccion('A3. PF-04, PF-05 Y OE5')
t51 = filas(tabla_con_epigrafe('Tabla 5.1:'))
fila = {f[0]: f[-1] for f in t51}
check('A3a Tabla 5.1: PF-04 no aprobada', fila.get('PF-04', '').startswith('NO APROBADA'), fila.get('PF-04', '')[:40])
check('A3b Tabla 5.1: PF-05 parcial', fila.get('PF-05', '').startswith('PARCIAL'), fila.get('PF-05', '')[:40])
check('A3c sin «PASS CON RESERVA» en todo el documento', 'PASS CON RESERVA' not in TODO)
check('A3d Tabla 5.1 con veredictos en castellano', all(fila.get('PF-0%d' % k, '').startswith('APROBADA') for k in (1, 2, 3)))
check('A3e §5.1.1 explica la reclasificación', 'no se aprobó' in b511 and 'en parte' in b511)
t61 = {f[0]: f for f in filas(tabla_con_epigrafe('Tabla 6.1:'))}
check('A3f Tabla 6.1: OE5 cumplido en parte', t61.get('OE5', [''])[-1].startswith('CUMPLIDO EN PARTE'),
      t61.get('OE5', [''])[-1])
check('A3g §6.2 ya no afirma que los cinco objetivos se cumplieron',
      'Los cinco objetivos específicos fueron cumplidos' not in TODO and 'OE5 se cumplió solo en parte' in b62)

# ============================================================================
seccion('A4. §6.3 Y §5.4: CIFRA MAL ATRIBUIDA Y ALCANCE DEL HALLAZGO')
check('A4a §6.3 atribuye 64,6 % a quitar reglas y ejemplos y 39,6 % a quitar también la base',
      'cae de 93,8 % a 39,6 % cuando faltan ambas cosas' not in TODO and '64,6 %' in b63 and '39,6 %' in b63)
check('A4b §6.3 sin «hallazgo más transferible» y acotado a la configuración estudiada',
      'hallazgo más transferible' not in TODO and 'configuración estudiada' in b63)
check('A4c §6.3 discute la regla crítica 5 y la sensibilidad al formato',
      'regla crítica 5' in b63 and '86,0 %' in b63 and 'Sclar et al., 2024' in b63)
check('A4d §6.3 presenta la conclusión como hipótesis a replicar', 'hipótesis a replicar' in b63)
check('A4e §5.4 no presenta el 74,0 % como referencia de lo alcanzable',
      'referencia documentada del desempeño alcanzable' not in TODO and '86,0 %' in b54)

# ============================================================================
seccion('A5. CRITERIOS FIJADOS ANTES O DESPUÉS DE MEDIR')
t511 = filas(tabla_con_epigrafe('Tabla 5.11:'))
check('A5a Tabla 5.11 tiene la columna del momento en que se fijó el criterio',
      bool(t511) and len(t511[0]) == 5 and 'se fijó' in t511[0][2], str(t511[:1]))
col = {f[0][:3]: f[2] for f in t511[1:]} if t511 and len(t511[0]) == 5 else {}
check('A5b H1: umbral de 10× después; 30 s antes', 'después' in col.get('H1:', '') and 'antes' in col.get('H1:', ''))
check('A5c H2a: extensión a cuatro categorías después', 'después' in col.get('H2a', ''))
check('A5d H2b: regla fijada antes del diseño factorial', 'antes de medir en el diseño factorial' in col.get('H2b', ''))
check('A5e §5.3 declara qué contraste es prospectivo y cita la versión previa',
      'no constituye una confirmación' in b53 and '9dc7b44' in b53)
check('A5f §6.1 matiza la confirmación', 'sin valor de confirmación' in b61)
check('A5g el resumen matiza la confirmación', 'regla fijada antes de medir' in RES)
check('A5h el abstract matiza la confirmación', 'rule fixed before measuring' in ABS)
nres, nabs = len(RES.split()), len(ABS.split())
check('A5i resumen dentro de 300 palabras', 250 <= nres <= 300, 'tiene %d' % nres)

# ============================================================================
seccion('A6. BASELINE: ACUERDO ENTRE INSTRUMENTOS Y SESGO DE EXPECTATIVA')
check('A6a sin «instrumento independiente», «validado de forma independiente» ni «vías independientes»',
      not re.search(r'instrumentos? independientes?|validado de forma independiente|vías independientes', TODO))
check('A6b sin «triangulación» en el método', 'triangula' not in bloque('3.2', '3.3'))
check('A6c la fórmula pedida aparece en §3.2 y §5.1.4',
      'acuerdo entre dos instrumentos sobre una misma sesión' in bloque('3.2', '3.3')
      and 'acuerdo entre dos instrumentos sobre una misma sesión' in b514)
check('A6d §3.6 nombra el sesgo de expectativa del operador', 'sesgo de expectativa' in b36.lower())

# ============================================================================
seccion('A7. PRECISIONES DEL DISEÑO FACTORIAL')
check('A7a §3.5.6 define la métrica de similitud', 'SequenceMatcher' in b356)
check('A7b §5.2.4 informa la exactitud media por repetición con intervalo',
      '93,1 %' in b524 and '76,7 %' in b524 and re.search(r'media por repetición[^.]*IC 95 %', b524, re.I) is not None)
t510 = filas(tabla_con_epigrafe('Tabla 5.10:'))
check('A7c Tabla 5.10 tiene la fila de la media por repetición',
      any(f and f[0].startswith('Media por repetición') for f in t510))
check('A7d §3.5.3 describe las definiciones que recibió el evaluador independiente',
      'Si dudás entre dos' in b353 or 'la que más pese' in b353)
check('A7e §3.5.3 aclara que no recibió las reglas del prompt', 'no recibió las reglas' in b353)

# ============================================================================
seccion('B. INCONSISTENCIAS REMANENTES DE LA TABLA FINAL DEL DICTAMEN')
t31 = filas(tabla_con_epigrafe('Tabla 3.1:'))
fila_f2 = next((f for f in t31 if f and f[0].startswith('4.')), [''])
check('B1  Tabla 3.1 no cuenta el trigger de Gmail como nodo en uso',
      '18 nodos funcionales' not in ' '.join(fila_f2) and 'Gmail' in ' '.join(fila_f2), ' '.join(fila_f2)[:80])
check('B2  «casi simultáneas» en todo el documento', not re.search(r'(?<!casi )simultáneas', TODO))
cap510 = next((p for p in P if p.startswith('Tabla 5.10:')), '')
check('B3  el título de la Tabla 5.10 distingue exhaustividad y exactitud', 'xhaustividad' in cap510, cap510)
check('B4  citas con «y», no con «&»',
      not re.search(r'[A-ZÁÉÍÓÚ][\wáéíóúñ]+ & [A-ZÁÉÍÓÚ][\wáéíóúñ]+,? \(?\d{4}', CAP1_7 + '\n' + CELDAS))
amen = [p for p in d.paragraphs if p.text.strip().startswith(('Amenaza:', 'Mitigación:'))]
todo_cursiva = [p.text[:30] for p in amen if len(p.runs) and all(r.italic for r in p.runs if r.text.strip())]
check('B5  §3.6: la cursiva alcanza solo al rótulo', not todo_cursiva, str(todo_cursiva))
epig = [p for p in d.paragraphs if re.match(r'^Tabla [\dA-K]+\.\d+:', p.text.strip())]
sueltos = [p.text[:12] for p in epig if not p.paragraph_format.keep_with_next]
check('B6  todo epígrafe de tabla se mantiene en la página de su tabla', len(epig) >= 29 and not sueltos, str(sueltos))

# ============================================================================
seccion('C. ORIGEN DEL CORPUS Y VERSIÓN DEL REPOSITORIO')
check('C1  ya no se afirma que el equipo escribió los 150 mensajes',
      'y los 150 mensajes del corpus' not in TODO and 'construido por el propio equipo' not in TODO)
check('C2  §3.5.3 declara la redacción asistida por un modelo de lenguaje', 'Claude' in b353 and 'modelo de lenguaje' in b353)
check('C3  §3.6.2 lo trata como amenaza', 'más regular' in bloque('3.6.2', '3.6.3'))
check('C4  la Declaración de originalidad lo declara', 'corpus' in bloque('DECLARACIÓN', 'RESUMEN'))
check('C5  la etiqueta de la versión citada es la nueva', TODO.count('entrega-2026-09-r2') == 2
      and not re.search(r'entrega-2026-09(?!-r2)', TODO))

print()
print('=' * 78)
if fallos:
    print(' RESULTADO: %d/%d — %d FALLAS' % (n_ok, n_ok + len(fallos), len(fallos)))
    print('=' * 78)
    sys.exit(1)
print(' DICTAMEN DEL 15/09: %d/%d' % (n_ok, n_ok))
print('=' * 78)
