# -*- coding: utf-8 -*-
"""Auditoría de cuarta instancia (18/09/2026): los siete hallazgos, uno o más controles cada uno.

Se escribió antes de corregir: la corrida inicial falla en todos.
"""
import os
import re
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

RUTA = 'docs/TESIS_FINAL_UTN_v6.docx'
d = Document(RUTA)
P = [p.text.strip() for p in d.paragraphs]
TODO = '\n'.join(P) + '\n' + '\n'.join(
    c.text.strip() for t in d.tables for r in t.rows for c in r.cells)
fallos = []
n_ok = 0


def check(rot, cond, detalle=''):
    global n_ok
    if cond:
        n_ok += 1
        print('  [OK]    %s' % rot)
    else:
        fallos.append(rot)
        print('  [FALLA] %s %s' % (rot, detalle))


def seccion(titulo):
    est = [(i, p.style.name, p.text.strip()) for i, p in enumerate(d.paragraphs)]
    ini = [i for i, e, t in est if t == titulo and e.startswith('Heading')]
    if len(ini) != 1:
        return ''
    sig = [i for i, e, t in est if i > ini[0] and e.startswith('Heading')]
    return '\n'.join(P[ini[0] + 1:sig[0] if sig else len(P)])


def filas(*enc):
    t = [t for t in d.tables if [c.text.strip() for c in t.rows[0].cells][:len(enc)] == list(enc)]
    return [[c.text.strip() for c in r.cells] for r in t[0].rows] if len(t) == 1 else []


def primera_col(tabla_filas):
    return [f[0] for f in tabla_filas]


print('== M-01 · PF-06 y PC-06 en el inventario de pruebas')
tablas_pruebas = [[[c.text.strip() for c in r.cells] for r in t.rows] for t in d.tables
                  if [c.text.strip() for c in t.rows[0].cells][:4] == ['#', 'Prueba', 'Entrada', 'Resultado esperado']]
ids = [f[0] for t in tablas_pruebas for f in t]
check('M-01b las Tablas 3.4 y 3.5 incluyen PF-06 y PC-06', 'PF-06' in ids and 'PC-06' in ids, str(ids))
t51 = filas('Prueba', 'Escenario', 'Resultado')
check('M-01c la Tabla 5.1 informa PF-06', 'PF-06' in primera_col(t51), str(primera_col(t51)))
check('M-01d la §5.2.1 informa PC-06', 'PC-06' in seccion('5.2.1 Pruebas funcionales del chatbot'))
check('M-01e el recuento de pruebas funcionales pasa a doce',
      '10 pruebas funcionales' not in TODO and 'diez pruebas funcionales' not in TODO
      and TODO.count('12 pruebas funcionales') + TODO.count('doce pruebas funcionales') >= 3)
check('M-01f se declara que PF-06 y PC-06 son posteriores a las mediciones',
      'posteriores a las mediciones' in seccion('3.5.8 Diseño de las pruebas funcionales, de carga y de concurrencia'))

print()
print('== M-02 · la comprobación del camino completo, citada')
j = [p for p in P if p.startswith('(j) Construcción de las consultas')]
check('M-02a la §5.5 (j) cita la comprobación del camino completo',
      bool(j) and 'verificar_camino_completo' in j[0])
anexo_l = seccion('Anexo L: Registro de desvíos y correcciones')
check('M-02b el Anexo L la cita con lo que verificó y su limpieza',
      'verificar_camino_completo' in anexo_l and 'no integra' in anexo_l)

print()
print('== B-01 a B-05')
check('B-01 PF-06 y PC-06 se distinguen en la §5.5 (j) y en el Anexo L',
      bool(j) and 'PC-06' in j[0] and 'PC-06' in anexo_l)
ev = [f for f in os.listdir('experiments/PF/resultados') if f.startswith('pf06')]
txt = open('experiments/PF/resultados/' + ev[0], encoding='utf-8').read() if ev else ''
check('B-02 la evidencia de PF-06 identifica la versión publicada del workflow', 'versionId' in txt or 'versión publicada' in txt)
check('B-03 el documento cita la etiqueta r6', TODO.count('entrega-2026-09-r6') == 2
      and not re.search(r'entrega-2026-09(?!-r6)', TODO))
check('B-04 el recuento de vistas es uniforme', 'las cinco vistas de métricas' not in TODO
      and 'Las cinco vistas de métricas' not in TODO)
s351 = seccion('3.5.1 Recolección automatizada de métricas')
check('B-05a la §3.5.1 describe cómo se selecciona cada corrida', 'ORD-E1A-' in s351 and 'E8-' in s351)
s433 = seccion('4.3.3 Métricas MTTD y MTTR')
check('B-05b la §4.3.3 distingue la vista que define la métrica de la consulta que produjo cada valor',
      'define' in s433 and 'analizar_e1.sql' in s433)

print()
print('=' * 78)
if fallos:
    print(' RESULTADO: %d/%d — %d FALLAS' % (n_ok, n_ok + len(fallos), len(fallos)))
    print('=' * 78)
    sys.exit(1)
print(' AUDITORÍA 4.ª INSTANCIA: %d/%d' % (n_ok, n_ok))
print('=' * 78)
