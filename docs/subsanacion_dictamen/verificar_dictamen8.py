# -*- coding: utf-8 -*-
"""Dictamen del 15/09, grupo B (recomendadas) y observaciones de sección no incluidas en el grupo A.

Un control por pedido. Se escribió antes de corregir.
"""
import sys
import io
import re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

d = Document('docs/TESIS_FINAL_UTN_v6.docx')
P = [p.text.strip() for p in d.paragraphs]
HEAD = [p.text.strip() for p in d.paragraphs if p.style.name.startswith('Heading') and p.text.strip()]
TXT_P = '\n'.join(P)
CELDAS = '\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
TODO = TXT_P + '\n' + CELDAS
i_ref = next(k for k, t in enumerate(P) if t.startswith('CAPÍTULO 8'))
i_anx = next(k for k, t in enumerate(P) if t.startswith('CAPÍTULO 9'))
CAP1_7 = '\n'.join(P[:i_ref])
REFS = [t for t in P[i_ref + 1:i_anx] if t]

fallos = []
n_ok = 0


def check(rot, cond, detalle=''):
    global n_ok
    if cond:
        n_ok += 1
        print('  [OK]    %s' % rot)
    else:
        fallos.append(rot + ((' — ' + detalle) if detalle else ''))
        print('  [FALLA] %s %s' % (rot, detalle))


def bloque(desde, hasta):
    try:
        i = next(k for k, t in enumerate(P) if t.startswith(desde))
        j = next(k for k, t in enumerate(P) if k > i and t.startswith(hasta))
    except StopIteration:
        return ''
    return '\n'.join(P[i:j])


def tabla_con_epigrafe(prefijo):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefijo):
            sig = p._p.getnext()
            for t in d.tables:
                if t._tbl is sig:
                    return t
    return None


def filas(t):
    return [[c.text.strip() for c in r.cells] for r in t.rows] if t is not None else []


def seccion(titulo):
    print()
    print('=' * 78)
    print(' ' + titulo)
    print('=' * 78)


i_L = next((k for k, t in enumerate(P) if t.startswith('Anexo L:')), None)
ANX_L = '\n'.join(P[i_L:]) if i_L is not None else ''
b142 = bloque('1.4.2', '1.5 ')
b152 = bloque('1.5.2', '1.6 ')
b32 = bloque('3.2 ', '3.3 ')
b24 = bloque('2.4 Estado del arte', '2.5 ')
b25 = bloque('2.5 Tratamiento', 'CAPÍTULO 3')
b521 = bloque('5.2.1', '5.2.2')
b54 = bloque('5.4 Análisis', '5.5 ')   # 17/09 (revisión de estructura): 5.4.1 pasó a ser 5.5
b63 = bloque('6.3', '6.4')
b64 = bloque('6.4', 'CAPÍTULO 7')
b71 = bloque('7.1', '7.2')
b72 = bloque('7.2', 'CAPÍTULO 8')

# ============================================================================
seccion('B1. REGISTRO DE REVISIONES EN UN ANEXO')
check('B1a existe el Anexo L de desvíos y correcciones', i_L is not None and 'desvíos y correcciones' in P[i_L])
check('B1b el cuerpo no narra la formulación original de las hipótesis', 'formulación original' not in CAP1_7)
check('B1c el Anexo L recoge las formulaciones originales y la ablación E7',
      'formulación original' in ANX_L and 'cinco de los ocho bloques' in ANX_L and 'régimen zero-shot' in ANX_L)
check('B1d §3.5.6 y §5.2.4 no narran la ablación descartada',
      'cinco de los ocho bloques' not in CAP1_7 and 'Los 6,7 puntos que E7 atribuyó' not in CAP1_7
      and 'corrige la interpretación que el trabajo había extraído' not in CAP1_7)
check('B1e §3.6.4 y §4.4.3 remiten al anexo en lugar de narrar el error de auditoría',
      'La auditoría fue correcta en su método' not in CAP1_7 and 'La auditoría fue correcta en su método' in ANX_L
      and 'cuatro minutos antes de la corrida' in ANX_L)
check('B1f la cita corregida del commit quedó en el anexo', 'f297c9e' in ANX_L and 'f297c9e' not in bloque('Anexo H', 'Anexo I'))
check('B1g §6.3 no narra la corrección de E7', 'corrige la que el trabajo había extraído' not in TODO)

# ============================================================================
seccion('B2. EXTENSIÓN Y REDUNDANCIA')
check('B2a §5.3 sin la precisión repetida sobre H2b por subcategoría', 'cabe una precisión metodológica adicional' not in TODO)
check('B2b §5.4 sin la propagación simplificada del intervalo', '687× a 872×' not in TODO)
check('B2c §6.4 enumera las limitaciones en forma breve', len(b64) < 2000, 'tiene %d caracteres' % len(b64))
check('B2d §6.1 remite a §5.4 sin repetir la salvedad del factor',
      'explica por qué ese factor se lee como una diferencia de casi tres órdenes' not in TODO)

# ============================================================================
seccion('B3. HIPÓTESIS Y OBJETIVOS')
check('B3a H1 se formula como objetivo de estimación', 'objetivo de estimación' in b142)
check('B3b OE1 incluye la comparación con el procesamiento manual', 'procesamiento manual' in next((p for p in P if p.startswith('Implementar un pipeline de procesamiento')), ''))
check('B3c §1.2 declara que el problema de la atención fragmentada se aborda en parte',
      'solo en parte' in next((p for p in P if p.startswith('Atención al cliente fragmentada')), ''))
t61 = {f[0]: f for f in filas(tabla_con_epigrafe('Tabla 6.1:'))}
check('B3d Tabla 6.1: OE1 informa la reducción respecto del proceso manual', '780×' in ' '.join(t61.get('OE1', [])))

# ============================================================================
seccion('B4. ESTADO DEL ARTE, MARCO LEGAL Y BIBLIOGRAFÍA')
check('B4a antecedentes empíricos sobre n8n', 'Amir y Atif (2026)' in b24 and 'Tang et al. (2026)' in b24)
check('B4b las dos referencias nuevas están en la lista',
      any(r.startswith('Amir, A. R., & Atif, S. M. (2026)') for r in REFS) and any(r.startswith('Tang, Y., Zhou, Y., & Chen, H. (2026)') for r in REFS))
check('B4c Moffatt pasa al marco legal', 'Moffatt' not in bloque('2.4.1', '2.4.2') and 'Columbia Británica' in b25)
check('B4d Sclar et al. dialoga con el hallazgo de formato en el estado del arte', 'Sclar et al. (2024)' in b24)
check('B4e la afirmación sobre código cerrado no descansa en el fabricante', 'código cerrado' not in TODO)
check('B4f sin «fulfillment automation» atribuido a Turban', 'fulfillment automation' not in TODO)
t21 = filas(tabla_con_epigrafe('Tabla 2.1:'))
check('B4g Tabla 2.1 suma los dos antecedentes', any(f[0].startswith('Amir y Atif') for f in t21) and any(f[0].startswith('Tang et al.') for f in t21))

# ============================================================================
seccion('B5. ESCALA DE MADUREZ Y ENCUADRE METODOLÓGICO')
check('B5a sin la escala de madurez propia', 'escala de madurez' not in TODO)
check('B5b §3.2 justifica el encuadre con razones metodológicas',
      'estructurado como estudio de caso' not in TODO and 'unidades de análisis incrustadas' in b32)
check('B5c el diseño factorial se presenta como experimento controlado dentro del caso', 'experimento controlado' in b32)

# ============================================================================
seccion('B6. PROMPT, MODELO Y LÍNEAS FUTURAS')
check('B6a §7.1 recomienda llevar el mensaje del cliente al rol de usuario', 'mensaje del usuario' in b71)
check('B6b §7.2 recomienda fijar temperatura y versión fechada', 'temperatura' in b72 and 'versión fechada' in b72)
check('B6c sin líneas genéricas (Redis, sentimiento, ERP/CRM)',
      not re.search(r'\bRedis\b', TODO) and 'Análisis de sentimiento' not in TODO and 'ERP/CRM' not in TODO)

# ============================================================================
seccion('B7. PRUEBAS DEL FLUJO 2 CON EVIDENCIA')
check('B7a §5.2.1 remite a la evidencia versionada', 'experiments/PF' in b521 and '17 de septiembre' in b521)
check('B7b §5.2.1 informa PC-05 como no aprobada', 'PC-05 no se aprobó' in b521)
t35 = {f[0]: f for f in filas(tabla_con_epigrafe('Tabla 3.5:'))}
check('B7c Tabla 3.5: PC-02 usa un pedido que existe', 'ORD-HIST-001' in ' '.join(t35.get('PC-02', [])))
check('B7d Tabla 6.1: OE5 cuenta PC-05', 'PC-05' in ' '.join(t61.get('OE5', [])))
check('B7e §5.2.1 declara los tickets sin vínculo con la interacción', 'clave foránea' in b521)

# ============================================================================
seccion('B8. VERSIÓN')
check('B8a el documento cita la etiqueta entrega-2026-09-r6', TODO.count('entrega-2026-09-r6') == 2
      and not re.search(r'entrega-2026-09(?!-r6)', TODO))

print()
print('=' * 78)
if fallos:
    print(' RESULTADO: %d/%d — %d FALLAS' % (n_ok, n_ok + len(fallos), len(fallos)))
    print('=' * 78)
    sys.exit(1)
print(' DICTAMEN DEL 15/09, GRUPO B: %d/%d' % (n_ok, n_ok))
print('=' * 78)
