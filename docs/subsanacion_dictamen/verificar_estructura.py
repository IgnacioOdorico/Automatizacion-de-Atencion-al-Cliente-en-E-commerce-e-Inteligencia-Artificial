# -*- coding: utf-8 -*-
"""Revisión de estructura del documento (17/09/2026).

Controla lo que un evaluador ve al hojear: numeración de páginas, portada e índice
en su lugar, jerarquía de títulos coherente, títulos con el mismo criterio de
mayúsculas, sin separadores sueltos, sin secciones con una única subsección, y
listados de figuras y tablas que coinciden con sus epígrafes y con la sección en la
que cada una aparece. Si existe el PDF, controla además que cada título de primer
nivel empiece página y que el índice impreso coincida con las páginas reales.
Se escribió antes de corregir.
"""
import os
import re
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from lxml import etree
from docx import Document

RUTA = 'docs/TESIS_FINAL_UTN_v6.docx'
d = Document(RUTA)
P = d.paragraphs
body = list(d.element.body.iterchildren())
fallos = []
n_ok = 0


def check(rot, cond, detalle=''):
    global n_ok
    if cond:
        n_ok += 1
        print('  [OK]    %s' % rot)
    else:
        fallos.append(rot + ((' — ' + detalle) if detalle else ''))
        print('  [FALLA] %s %s' % (rot, detalle))


def xml(el):
    return etree.tostring(el).decode('utf-8')


# ------------------------------------------------------------------ páginas
s = d.sections[0]
pie = xml(s.footer._element)
check('S1  el pie de página numera las páginas', 'PAGE' in pie and 'fldChar' in pie)
check('S2  el encabezado no tiene caracteres sueltos', not any(p.text.strip() for p in s.header.paragraphs))
check('S3  la portada no lleva número', s.different_first_page_header_footer
      and not any(p.text.strip() for p in s.first_page_footer.paragraphs))

# ------------------------------------------------------------------ portada
i_decl = next(i for i, p in enumerate(P) if p.text.strip() == 'DECLARACIÓN DE ORIGINALIDAD')
portada = [p.text.strip() for p in P[:i_decl] if p.text.strip()]
repetidas = [t for t in portada if t == 'PORTADA' or t.startswith(('Título:', 'Autores:')) or t == 'Mendoza – Argentina, 2026']
check('S16 portada sin rótulos internos ni datos repetidos', not repetidas, str(repetidas))

# ------------------------------------------------------------------ parte inicial
check('S4  sin reglas horizontales decorativas', not any('o:hr' in xml(el) for el in body),
      str([i for i, el in enumerate(body) if 'o:hr' in xml(el)]))
indice = [el for el in body if el.tag.endswith('}sdt') and 'ÍNDICE' in ''.join(el.itertext())[:20]]
check('S5  el índice empieza en página nueva', bool(indice) and 'pageBreakBefore' in xml(indice[0]).split('</w:p>')[0])
inicial = {p.text.strip(): p for p in P if p.text.strip() in ('Listado de figuras', 'Listado de tablas', 'Glosario de siglas y acrónimos')}
check('S6  listados y glosario como títulos de primer nivel, en página nueva',
      len(inicial) == 3 and all(p.style.name == 'Heading 1' and p.paragraph_format.page_break_before for p in inicial.values()),
      str([(t, p.style.name) for t, p in inicial.items()]))

# ------------------------------------------------------------------ títulos
heads = [(p.style.name, p.text.strip()) for p in P if p.style.name.startswith('Heading') and p.text.strip()]
nums = [re.match(r'^(\d+(?:\.\d+)*)\s', t).group(1) for _, t in heads if re.match(r'^\d+(?:\.\d+)*\s', t)]
unicas = [n for n in nums if len([m for m in nums if m.startswith(n + '.') and m.count('.') == n.count('.') + 1]) == 1]
check('S7  ninguna sección tiene una única subsección', not unicas, str(unicas))
nivel_ok = all((st == 'Heading %d' % (t.split()[0].count('.') + 1)) for st, t in heads if re.match(r'^\d+\.\d+', t))
check('S8  el nivel de cada título corresponde a su numeración', nivel_ok,
      str([(st, t[:30]) for st, t in heads if re.match(r'^\d+\.\d+', t) and st != 'Heading %d' % (t.split()[0].count('.') + 1)]))
PROPIOS = r'^(PostgreSQL|Docker|Compose|Grafana|GPT|MTTD|MTTR|TMR|IA|LLMs?|WhatsApp|Telegram|Flujo|Chatbot|SERVQUAL|FAQ|n8n)'


def mayusculas_de_titulo(t):
    palabras = re.split(r'\s+', re.sub(r'^(\d+(?:\.\d+)*|Anexo [A-L]:)\s+', '', t))
    # la primera palabra y la que sigue a «—» (un subtítulo) llevan mayúscula
    return [w for previa, w in zip(palabras, palabras[1:])
            if w[:1].isupper() and not w.isupper() and not re.match(PROPIOS, w) and previa != '—']


titulo = [t for _, t in heads if re.match(r'^(\d+(?:\.\d+)*|Anexo [A-L]:)\s', t) and mayusculas_de_titulo(t)]
check('S9  títulos con mayúscula solo inicial y en nombres propios', not titulo, str(titulo))
secuencia = []
for _, t in heads:
    m = re.match(r'^(\d+)(?:\.(\d+))?(?:\.(\d+))?\s', t)
    if m:
        secuencia.append(tuple(int(x) for x in m.groups() if x))
saltos = []
for a, b in zip(secuencia, secuencia[1:]):
    ok = (len(b) == len(a) + 1 and b[:-1] == a and b[-1] == 1) or \
         (len(b) <= len(a) and b[:-1] == a[:len(b) - 1] and b[-1] == a[len(b) - 1] + 1) or \
         (b == (a[0] + 1, 1))   # primera sección del capítulo siguiente
    if not ok:
        saltos.append((a, b))
check('S10 numeración de secciones correlativa', not saltos, str(saltos[:5]))
anexos = [t for _, t in heads if t.startswith('Anexo ')]
check('S11 anexos correlativos de A a L', [a.split(':')[0][-1] for a in anexos] == list('ABCDEFGHIJKL'), str(anexos))


# ------------------------------------------------------------------ listados
def seccion_de(idx):
    s_ = None
    for p in P[:idx]:
        if p.style.name.startswith('Heading'):
            m = re.match(r'^(\d+(?:\.\d+)*)\s', p.text.strip())
            if m:
                s_ = m.group(1)
            elif p.text.strip().startswith('Anexo'):
                s_ = p.text.strip().split(':')[0]
    return s_


caps = {}
for i, p in enumerate(P):
    m = re.match(r'^(Figura \d+|Tabla [\dA-K]+\.\d+): (.*)$', p.text.strip())
    if m:
        caps[m.group(1)] = (m.group(2).rstrip('.'), seccion_de(i))
problemas = []
for t in d.tables:
    h = [c.text.strip() for c in t.rows[0].cells]
    if len(h) == 3 and h[0] in ('Figura', 'Tabla') and h[2] == 'Sección':
        for r in t.rows[1:]:
            k, desc, sec = [c.text.strip() for c in r.cells]
            if k not in caps:
                problemas.append('%s sin epígrafe' % k)
                continue
            cap, s_ = caps[k]
            if not cap.startswith(desc.rstrip('.')):
                problemas.append('%s: descripción distinta del epígrafe' % k)
            if s_ != sec:
                problemas.append('%s: sección %s en el listado, %s en el documento' % (k, sec, s_))
check('S12 listados de figuras y tablas coinciden con epígrafes y secciones', not problemas, str(problemas))
figs = [int(k.split()[1]) for k in caps if k.startswith('Figura')]
check('S13 figuras numeradas de 1 a n sin huecos', sorted(figs) == list(range(1, len(figs) + 1)), str(sorted(figs)))

# ------------------------------------------------------------------ PDF
pdf = RUTA.replace('.docx', '.pdf')
if os.path.exists(pdf) and os.path.getmtime(pdf) >= os.path.getmtime(RUTA) - 120:
    import fitz
    doc = fitz.open(pdf)
    no_empieza = []
    for i, pg in enumerate(doc):
        bl = sorted([b for b in pg.get_text('blocks') if b[4].strip() and not re.match(r'^(\d+|\|)$', b[4].strip())],
                    key=lambda b: b[1])
        for j, b in enumerate(bl):
            t = b[4].strip()
            if re.match(r'^(CAPÍTULO \d+: [^.]*$|RESUMEN$|ABSTRACT$|DECLARACIÓN DE ORIGINALIDAD$|ÍNDICE$|Listado de (figuras|tablas)$|Glosario de siglas y acrónimos$)', t) and j > 0:
                no_empieza.append('p.%d %s' % (i + 1, t[:30]))
    check('S14 (PDF) cada título de primer nivel empieza página', not no_empieza, str(no_empieza))
    numeradas = sum(1 for pg in doc[1:] if re.search(r'^\s*%d\s*$' % (pg.number + 1), pg.get_text(), re.M))
    check('S15 (PDF) las páginas llevan su número impreso', numeradas >= doc.page_count - 2, '%d de %d' % (numeradas, doc.page_count))
else:
    print('  (PDF desactualizado: se omiten S14 y S15)')

print()
print('=' * 78)
if fallos:
    print(' RESULTADO: %d/%d — %d FALLAS' % (n_ok, n_ok + len(fallos), len(fallos)))
    print('=' * 78)
    sys.exit(1)
print(' ESTRUCTURA: %d/%d' % (n_ok, n_ok))
print('=' * 78)
