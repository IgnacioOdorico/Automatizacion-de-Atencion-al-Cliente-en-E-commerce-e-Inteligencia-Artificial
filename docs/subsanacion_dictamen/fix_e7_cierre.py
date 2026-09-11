# -*- coding: utf-8 -*-
"""E7, cierre: listado de tablas, resumen/abstract y el residuo de (a-bis).

El apartado (a-bis) de §5.4.1 proponia como mitigacion futura "incorporar al
prompt ejemplos etiquetados de esta clase (few-shot), que hoy no tiene ninguno".
Con la configuracion medida establecida, eso es falso: el prompt tiene siete
ejemplos. La mitigacion que corresponde es otra.
"""
import sys
import os
import copy
sys.path.insert(0, 'docs/subsanacion_dictamen')
from docxkit import *
from docx import Document
from docx.oxml.ns import qn

if any(f.startswith('~$') for f in os.listdir('docs')):
    sys.exit('ERROR: Word tiene abierto un documento en docs/. Cerralo primero.')

RUTA = 'docs/TESIS_FINAL_UTN_v6.docx'
d = Document(RUTA)


def par(pref):
    for p in d.paragraphs:
        if p.text.strip().startswith(pref):
            return p
    raise KeyError(pref)


print('=' * 78)
print(' 1. Listado de Tablas: falta la 5.10')
print('=' * 78)
for t in d.tables:
    if t.rows[0].cells[0].text.strip() == 'Tabla' and t.rows[0].cells[2].text.strip() == 'Sección':
        j = [i for i, r in enumerate(t.rows) if r.cells[0].text.strip() == 'Tabla 5.11'][0]
        tr = copy.deepcopy(t.rows[j]._tr)
        t.rows[j]._tr.addprevious(tr)
        k = [i for i, r in enumerate(t.rows) if r._tr is tr][0]
        set_cell(t, k, 0, 'Tabla 5.10')
        set_cell(t, k, 1, 'Exactitud de clasificación por clase, con y sin base de conocimiento')
        set_cell(t, k, 2, '5.2.4')
        print('  entrada de la Tabla 5.10 insertada antes de la 5.11')
        break

print()
print('=' * 78)
print(' 2. El resumen y el abstract reportan la ablacion')
print('=' * 78)
n = replace_everywhere(d,
    'y la precisión de clasificación alcanzó el 92,7 % (IC 95 % de Wilson: 87,3 % a 95,9 %).',
    'y la exactitud de clasificación alcanzó el 92,7 % (IC 95 % de Wilson: 87,3 % a 95,9 %) con '
    'un prompt que inyecta la base de conocimiento de la tienda y presenta siete ejemplos '
    'etiquetados. Una ablación sobre el mismo corpus y las mismas etiquetas de referencia, con un '
    'prompt reducido que no incluye ninguna de las dos cosas, arroja 86,0 % (IC 95 %: 79,5 % a '
    '90,7 %): una diferencia de 6,7 puntos que la prueba de McNemar sobre el diseño apareado '
    'confirma como no atribuible al azar (p = 0,031) y que se concentra en la categoría de '
    'consultas frecuentes, la que pierde 14,6 puntos.')
print('  resumen en castellano: %d' % n)

n = replace_everywhere(d,
    'and an intent classification accuracy of 92.7 % (Wilson 95 % CI: 87.3 % to 95.9 %) in a '
    'zero-shot regime.',
    'and an intent classification accuracy of 92.7 % (Wilson 95 % CI: 87.3 % to 95.9 %) with a '
    'prompt that injects the store knowledge base and presents seven labelled examples. An '
    'ablation over the same corpus and the same reference labels, with a reduced prompt including '
    'neither, yields 86.0 % (95 % CI: 79.5 % to 90.7 %): a 6.7-point difference that McNemar’s '
    'test on the paired design confirms as not attributable to chance (p = 0.031), concentrated on '
    'the frequently-asked-question category, which loses 14.6 points.')
print('  abstract en ingles: %d' % n)

print()
print('=' * 78)
print(' 3. Residuo en (a-bis): el prompt SI tiene ejemplos')
print('=' * 78)
n = replace_everywhere(d,
    'o en incorporar al prompt ejemplos etiquetados de esta clase (few-shot), que hoy no tiene '
    'ninguno, sin requerir fine-tuning del modelo.',
    'o en ampliar los siete ejemplos etiquetados que el prompt ya incluye con casos de frontera '
    'entre GENERAL y FAQ, que hoy no están representados entre ellos, sin requerir ajuste fino '
    'del modelo. La Sección 5.2.4 aporta evidencia de que esa frontera es justamente la más '
    'sensible al contenido del prompt.')
print('  (a-bis) corregida: %d' % n)

d.save(RUTA)
print()
print('guardado.')

# ================================ verificacion ================================
d2 = Document(RUTA)
P2 = [p.text.strip() for p in d2.paragraphs]
TODO = '\n'.join(P2) + '\n' + '\n'.join(
    c.text for t in d2.tables for r in t.rows for c in r.cells)
HEAD = [(p.style.name, p.text.strip()) for p in d2.paragraphs if p.style.name.startswith('Heading')]
fallos = []


def check(rot, cond, det=''):
    print('  [%s] %s %s' % ('OK   ' if cond else 'FALLA', rot, det))
    if not cond:
        fallos.append(rot)


print()
print('=== control del bloque E7 completo ===')
check('§2.2.2 declara las dos configuraciones',
      'La configuración principal opera en régimen few-shot con recuperación de contexto' in TODO)
check('§3.5.6 protocolo de la ablación',
      any(t.startswith('3.5.6 Ablación del contexto') for _, t in HEAD))
check('§3.5.6 declara McNemar y el diseño apareado',
      'la prueba de McNemar sobre los pares discordantes' in TODO)
check('§4.4.3 describe la configuración medida (6338)',
      'tiene 6338 caracteres y consta de cinco bloques' in TODO)
check('§4.4.3 declara el error de auditoría',
      'se auditó la versión vigente del workflow y no la que había producido la medición' in TODO)
check('§4.4.3 ya no afirma zero-shot',
      'no en régimen zero-shot' in TODO and 'El prompt no incorpora ejemplos etiquetados' not in TODO)
check('§5.2.4 con los resultados',
      any(t.startswith('5.2.4 Ablación') for _, t in HEAD))
check('Tabla 5.10 existe y tiene las dos condiciones',
      any(r.cells[0].text.strip() == 'FAQ' and '77,1' in ' '.join(c.text for c in r.cells)
          for t in d2.tables for r in t.rows))
check('Tabla 5.10 en el listado', 'Exactitud de clasificación por clase, con y sin base' in TODO)
check('la antigua 5.10 es ahora 5.11', 'Tabla 5.11: Contrastación de hipótesis' in TODO)
check('McNemar reportado en resultados', 'p exacto de la binomial bilateral de 0,031' in TODO
      or 'valor p exacto de la binomial bilateral de 0,031' in TODO)
check('H2b acotado a la condición', 'CONFIRMADA con base de conocimiento; NO CONFIRMADA sin ella' in TODO)
check('§5.4 sobre las dos condiciones', 'sin ninguna de las dos cosas, 86,0 %' in TODO)
check('§6.3 con el hallazgo reformulado',
      'lo que una PyME tiene que construir y mantener no es el prompt sino su base de preguntas' in TODO)
check('§7.2 sin la línea ya cumplida',
      'Incorporación de ejemplos etiquetados al prompt y cableado' not in TODO)
check('Anexo H publica los dos prompts',
      'Se transcriben los dos prompts de sistema que este trabajo midió' in TODO)
check('Anexo H remite al commit verificable', 'commit f297c9e' in TODO)
check('resumen reporta la ablación', 'arroja 86,0 %' in TODO)
check('abstract reporta la ablación', 'yields 86.0 %' in TODO)
check('(a-bis) ya no dice que no hay ejemplos',
      'que hoy no tiene ninguno' not in TODO)

print()
if fallos:
    print('FALLAS: %d' % len(fallos))
    for f in fallos:
        print('  - %s' % f)
    sys.exit(1)
print('BLOQUE E7 — SIN FALLAS')
