# -*- coding: utf-8 -*-
"""Dictamen del 14/09/2026 (7,2/10): una comprobación por observación.

Cubre las observaciones obligatorias 1 a 11 y las recomendadas que no exigen
reestructurar el documento. Cada control indica qué observación cierra. Se
escribió ANTES de corregir, de modo que su primera corrida tiene que fallar.
"""
import sys
import io
import re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

RUTA = 'docs/TESIS_FINAL_UTN_v6.docx'
C1_TXT = 'experiments/E8/prompts/C1.txt'
URL = 'https://github.com/IgnacioOdorico/Automatizacion-de-Atencion-al-Cliente-en-E-commerce-e-Inteligencia-Artificial'

d = Document(RUTA)
P = [p.text.strip() for p in d.paragraphs]
HEAD = [p.text.strip() for p in d.paragraphs if p.style.name.startswith('Heading')]
TXT_P = '\n'.join(P)
CELDAS = '\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
TODO = TXT_P + '\n' + CELDAS
i_ref = next(k for k, t in enumerate(P) if t.startswith('CAPÍTULO 8'))
i_anx = next(k for k, t in enumerate(P) if t.startswith('CAPÍTULO 9'))
CUERPO = '\n'.join(P[:i_ref]) + '\n' + '\n'.join(P[i_anx:])
REFS = [t for t in P[i_ref + 1:i_anx] if t]

fallos = []
n_ok = 0


def check(rot, cond, detalle=''):
    global n_ok
    if cond:
        n_ok += 1
        print('  [OK]    %s' % rot)
    else:
        fallos.append(rot + ((' — ' + detalle) if detalle else ''))
        print('  [FALLA] %s %s' % (rot, detalle))


def bloque(desde, hasta):
    try:
        i = next(k for k, t in enumerate(P) if t.startswith(desde))
        j = next(k for k, t in enumerate(P) if k > i and t.startswith(hasta))
    except StopIteration:
        return ''
    return '\n'.join(P[i:j])


def tabla_con_epigrafe(prefijo):
    """La tabla que sigue al epígrafe que empieza con `prefijo`."""
    for p in d.paragraphs:
        if p.text.strip().startswith(prefijo):
            sig = p._p.getnext()
            while sig is not None and sig.tag.endswith('}p') and not ''.join(sig.itertext()).strip():
                sig = sig.getnext()
            for t in d.tables:
                if t._tbl is sig:
                    return t
    return None


def filas(t):
    return [[c.text.strip() for c in r.cells] for r in t.rows] if t is not None else []


def seccion(titulo):
    print()
    print('=' * 78)
    print(' ' + titulo)
    print('=' * 78)


# ============================================================================
seccion('1-2. ABLACIÓN FACTORIAL (E8) Y VARIABILIDAD')
b356 = bloque('3.5.6', '3.5.7')
check('1a  §3.5.6 se titula como diseño factorial',
      '3.5.6 Diseño factorial del prompt: base de conocimiento, reglas y ejemplos' in HEAD)
# 17/09 (dictamen del 15/09, grupo B): la historia de E7 pasó al Anexo L.
check('1b  E7 quitaba cinco de los ocho bloques (Anexo L, con remisión desde §3.5.6)',
      'cinco de los ocho bloques' in TODO and 'Anexo L' in b356)
check('1c  §3.5.6 declara mayoría de tres repeticiones, McNemar y Holm',
      'exactitud por mayoría' in b356 and 'McNemar, 1947' in b356 and 'Holm (1979)' in b356)
check('2a  §3.5.6 declara la temperatura y las repeticiones',
      'valor por defecto de la API, que es 1' in b356 and 'tres veces' in b356)
b443 = bloque('4.4.3', '4.4.4')
check('1d  §4.4.3 describe los ocho bloques del prompt medido',
      'ocho bloques' in b443 and 'diez reglas críticas' in b443)
check('2b  §4.4.3 declara alias del modelo y temperatura', 'alias' in b443 and 'que es 1' in b443)
# 17/09 (dictamen del 15/09, grupo B): la reconstrucción pasó al Anexo L.
check('4a  §4.4.3 remite a la evidencia de trazabilidad',
      'experiments/E8/resultados/trazabilidad_versiones.txt' in TODO and 'f68da9d' in b443 and 'Anexo L' in b443)
b524 = bloque('5.2.4', '5.2.5')
check('1e  §5.2.4 con título nuevo',
      '5.2.4 Aporte de la base de conocimiento y de las reglas y ejemplos' in HEAD)
t510 = filas(tabla_con_epigrafe('Tabla 5.10:'))
check('1f  la Tabla 5.10 tiene las cuatro condiciones y la exactitud por mayoría',
      # 16/09: el dictamen del 15/09 pidió sumar la media por repetición, en una séptima fila
      len(t510) == 7 and t510[6][0].startswith('Media por repetición')
      and all(c in ' '.join(t510[0]) for c in ('C1', 'C2', 'C3', 'C4'))
      and '140 (93,3 %)' in ' '.join(t510[5]) and '111 (74,0 %)' in ' '.join(t510[5]), str(t510[:1])[:160])
check('1g  §5.2.4 reporta efectos, interacción, réplica y variabilidad',
      all(x in b524 for x in ('17,3 puntos', '2,0 puntos', '−5,3', 'p exacto = 1,0', '98,0 %')))
check('1h  §5.2.4 reporta las etiquetas fuera de vocabulario', '«CAMBIO»' in b524 and '«DEVOLUCION»' in b524)
check('1i  ninguna conclusión atribuye la clasificación a la base de conocimiento',
      'no es el prompt sino su base de preguntas frecuentes' not in TODO
      and 'la base de conocimiento es la que da el margen' not in TODO
      and 'lo que una PyME gana por mantener su base de conocimiento' not in TODO
      and 'es lo que hace que la categoría FAQ sea decidible' not in TODO)
check('1j  §6.3 invierte la conclusión práctica', 'no su base de preguntas frecuentes' in bloque('6.3', '6.4'))
bH = bloque('Anexo H', 'Anexo I')
c1 = io.open(C1_TXT, encoding='utf-8', newline='').read().replace('\r\n', '\n').strip()
fuente = [p.text.replace('\r', '').strip() for p in d.paragraphs if p.style.name == 'Source Code']
check('1k  el Anexo H transcribe íntegro el prompt medido (C1)', c1 in fuente)
check('1l  el Anexo H explica C1 a C4 y cita f68da9d',
      all(x in bH for x in ('C2 omite', 'C3 omite', 'C4 omite', 'commit f68da9d'))
      and 'los dos que siguen' not in TODO)
check('1m  f297c9e solo aparece como cita corregida',
      all('en un primer momento se citó' in p for p in P if 'f297c9e' in p))
check('1n  resumen y abstract reportan el factorial',
      '17,3' in bloque('RESUMEN', 'ABSTRACT') and '17.3' in bloque('ABSTRACT', 'Listado de Figuras'))

seccion('3. INDEPENDENCIA ENTRE PROMPT Y CORPUS')
b36 = bloque('3.6 Amenazas', 'CAPÍTULO 4')
check('3a  §3.6 declara la cronología prompt-corpus', 'ebbd17c' in b36 and '77a5a04' in b36 and '110 días' in b36)
check('3b  §3.6 declara la similitud y la sensibilidad', '0,89' in b36 and 'sensibilidad' in b36)

seccion('4. TRAZABILIDAD DEL ARTEFACTO MEDIDO')
b522 = bloque('5.2.2', '5.2.3')
check('4b  §5.2.2 no atribuye la diferencia a la red', 'por la latencia de red real' not in TODO)
check('4c  §5.2.2 declara que Telegram corrió con el prompt reducido', 'prompt reducido' in b522)
# 15/09: la discusión de la no comparabilidad entre canales pasó de §4.6 a §5.2.2.
check('4d  §5.2.2 cuenta el prompt entre los factores de no comparabilidad',
      'no son apareadas' in bloque('5.2.2', '5.2.3') and 'en el prompt con que corrieron' in bloque('5.2.2', '5.2.3'))
oe2 = [f for f in filas(tabla_con_epigrafe('Tabla 6.1:')) if f and f[0] == 'OE2']
check('4e  Tabla 6.1 OE2 dice sobre qué configuración se cumple',
      bool(oe2) and 'configuración vigente' in oe2[0][2] and '93,3 %' in oe2[0][2], str(oe2)[:120])

seccion('5. REPRODUCIBILIDAD')
check('5a  URL del repositorio en la Declaración', URL in bloque('DECLARACIÓN DE ORIGINALIDAD', 'RESUMEN'))
bJ = bloque('Anexo J', 'Anexo K')
check('5b  URL y versión en el Anexo J', URL in bJ and 'etiqueta' in bJ)
bE = bloque('Anexo E', 'Anexo F')
check('5c  el Anexo E remite a los guiones versionados de E1',
      'run_flujo1_carga.ps1' in bE and 'run_flujo1_concurrencia.ps1' in bE and 'ORD-CARGA-' not in bE)
tD = next((filas(t) for t in d.tables if filas(t) and filas(t)[0][:2] == ['Categoría', 'Pregunta']), [])
check('5d  el Anexo D transcribe íntegras las 23 entradas',
      len(tD) == 24 and not any(f[2].endswith('…') for f in tD[1:])
      and any('por 5 días hábiles' in f[2] for f in tD) and 'reproduce el inicio del texto' not in TODO)

seccion('6. PREGUNTA, HIPÓTESIS Y OBJETIVOS')
b14 = bloque('1.4.1', '1.5 Objetivos')
check('6a  la pregunta limita la reducción a las órdenes',
      'respecto del procesamiento manual' in b14 and 'con qué tiempo de respuesta' in b14)
preg = next((p for p in P if p.startswith('¿En qué medida')), '')
check('6i  §6.1 y el objetivo general citan la pregunta vigente',
      bool(preg) and preg in bloque('6.1', '6.2') and 'reduce los tiempos operativos del ciclo post-venta' not in TODO
      and 'con umbrales absolutos' in bloque('1.5.1', '1.5.2'))
h1 = next((p for p in P if p.startswith('H1:')), '')
# 17/09 (dictamen del 15/09, grupo B): H1 se presenta como objetivo de estimación.
check('6b  H1 operacionaliza la magnitud y declara cuándo',
      'objetivo de estimación' in h1 and 'orden de magnitud' in h1 and 'resultado conocido' in h1)
h2a = next((p for p in P if p.startswith('H2a:')), '')
check('6c  H2a coincide con el contraste (cuatro categorías)', 'cada una de las cuatro categorías' in h2a)
h2b = next((p for p in P if p.startswith('H2b:')), '')
check('6d  H2b fija la regla del límite inferior y cuándo se adoptó',
      'límite inferior del intervalo de confianza de Wilson' in h2b
      and 'con la corrida del corpus ya ejecutada' in h2b)
check('6e  no queda «H2» a secas', not re.search(r'H2(?![ab])', TODO),
      str(re.findall(r'.{30}H2(?![ab]).{10}', TODO))[:200])
check('6f  «precisión» no designa la exactitud global',
      not re.search(r'precisi[oó]n (de|en) clasificaci[oó]n', TODO, re.I)
      and 'Precisión (accuracy)' not in CELDAS.replace('rotulado «Precisión (accuracy)»', ''))
check('6g  se declaran los componentes incorporados durante la investigación',
      'se incorporaron durante la investigación' in bloque('1.5.2', '1.6 Alcance'))
tr = re.findall(r'.{0,40}tiempo real', TODO)
check('6h  «en tiempo real» solo donde corresponde',
      all('tracking' in x or 'seguir el paquete' in x for x in tr), str(tr)[:240])

seccion('7. ENCUADRE METODOLÓGICO')
b32 = bloque('3.2 Enfoque', '3.3 Metodología')
check('7a  el caso instrumental se atribuye a Stake (1995)',
      'Stake (1995)' in b32
      and 'Siguiendo la tipología de Yin (2018), se trata de un estudio de caso instrumental' not in TODO
      and any(r.startswith('Stake, R. E. (1995)') for r in REFS))
check('7b  se justifica Yin en un entorno simulado',
      'entorno de laboratorio' in b32 and 'Hevner et al. (2004)' in b32
      and any(r.startswith('Hevner, A. R.') for r in REFS))
check('7c  §3.3 reconoce el desarrollo iterativo', 'iterativo' in bloque('3.3 Metodología', '3.4 Herramientas'))

seccion('8. CONCLUSIONES SIN RESPALDO')
b61 = bloque('6.1', '6.2')
b63 = bloque('6.3', '6.4')
check('8a  §6.1 limita la reducción a las órdenes', 'Para el procesamiento de órdenes' in b61)
check('8b  no se afirma «tres hipótesis confirmadas»',
      'fueron confirmadas por los datos experimentales' not in TODO and 'CONFIRMADA' not in CELDAS)
check('8c  §6.3 no afirma «sin escribir código» y declara lo no medido',
      'sin escribir código de aplicación' not in TODO and 'cinco nodos Function' in b63 and 'no fue medido' in b63)
check('8d  §5.4 no afirma manejo de errores ortográficos sin datos',
      'manejando correctamente mensajes con errores ortográficos' not in TODO)
check('8e  §5.4.1 no habla de dominio «entrenado»', 'dominio diferente al entrenado' not in TODO)
check('8f  §6.3 separa validez estructural y semántica',
      'validez estructural' in b63 and 'resultó operativamente válida' not in TODO)
# 17/09 (dictamen del 15/09, grupo B): la escala propia se retiró; queda la distinción de Dumas et al.
check('8g  sin escala de madurez atribuida ni propia',
      'van der Aalst' not in TODO and 'escala de madurez' not in TODO and 'automatizar un proceso y gestionarlo' in TODO)

seccion('9. FACTOR 780×')
b54 = bloque('5.4 Análisis', '5.4.1')
check('9a  ningún lugar lee el factor como cota superior', 'cota superior' not in TODO and 'upper bound' not in TODO)
check('9b  los sesgos operan en direcciones opuestas', 'direcciones opuestas' in b54)
check('9c  significación práctica en tiempo absoluto', '41 minutos' in b54)

seccion('10. CONCURRENCIA')
b513 = bloque('5.1.3', '5.1.4')
check('10a pérdida silenciosa como falla de confiabilidad',
      'falla de confiabilidad' in b513 and 'HTTP 200' in b513 and 'seguro pero no elástico' not in TODO)
check('10b el CHECK (stock >= 0) explica la ausencia de sobreventa', 'CHECK (stock >= 0)' in b513)
check('10c sesgo de supervivencia en las latencias', 'sesgo de supervivencia' in b513)
check('10d concurrencia efectiva y seis rondas en dos ejecuciones',
      'concurrencia efectiva' in b513 and 'dos ejecuciones del guion' in b513)

seccion('11. INCONSISTENCIAS DEL CUADRO')
check('11a §4.5 cinco paneles heredan de la vista',
      'cinco lo heredan de la vista v_chatbot_corpus' in TODO and 'seis del tablero del Flujo 2 lo heredan' not in TODO)
check('11b §4.6.2 no dice que todas las vistas filtran',
      'Las vistas que alimentan los tableros de resultados filtran por ese valor' not in TODO)
check('11c §5.1.4 CV frente a r',
      'dispersión baja que indica que el procedimiento se ejecutó de manera estable' not in TODO
      and 'piso y no como un valor central' not in TODO)
check('11d Tabla 3.3 sin «Tasa de resolución autónoma»', 'Tasa de resolución autónoma' not in TODO)
check('11e §2.1.3 MTTR coherente con §4.3.3', 'hasta que el cliente recibe la notificación' not in TODO)
check('11f §2.4 criterio de preprints coherente',   # 15/09: el estado del arte pasó de §2.5 a §2.4
      'preprints de laboratorios o conferencias reconocidas' not in TODO
      and 'excepción declarada' in bloque('2.4 Estado del arte', '2.4.1'))
check('11g Tabla 5.3 y Figura 8 coherentes',
      '3 de verificación)' not in TODO and 'ORD-AUDIT-ALERT' in CELDAS
      and 'no integran el total de la Tabla 5.3' in TXT_P)
check('11h §5.1.1 sin «eventos»', '(emails, eventos)' not in TODO)
t49 = filas(tabla_con_epigrafe('Tabla 3.4:'))    # 15/09: las Tablas 4.9 y 4.10 pasaron a §3.5.8
t410 = filas(tabla_con_epigrafe('Tabla 3.5:'))
check('11i identificadores de pruebas unificados',
      [f[0] for f in t49[1:]] == ['PF-01', 'PF-02', 'PF-03', 'PF-04', 'PF-05']
      and [f[0] for f in t410[1:]] == ['PC-01', 'PC-02', 'PC-03', 'PC-04', 'PC-05'])
check('11j sin «e ejemplos»', not re.search(r'\be ejemplos', TODO))   # «siete ejemplos» no cuenta
check('11k §2.3.2 sin «varios órdenes de magnitud respecto al estándar manual»',
      'varios órdenes de magnitud respecto al estándar manual' not in TODO)
check('11l sin la alerta inexistente low_stock_alert', 'low_stock_alert' not in TODO)
b511 = bloque('5.1.1', '5.1.2')
t51 = filas(tabla_con_epigrafe('Tabla 5.1:'))
check('11m PF-04 y PF-05 con evidencia y mecanismo reales',
      'experiments/PF' in b511 and 'HTTP 200' in b511
      and not any('orden rechazada con error controlado' in ' '.join(f) for f in t51))

seccion('RECOMENDADAS')
i_res = P.index('RESUMEN')
res = ' '.join(P[i_res + 1:next(k for k, t in enumerate(P) if k > i_res and t.startswith('Palabras clave'))])
i_abs = P.index('ABSTRACT')
abs_ = ' '.join(P[i_abs + 1:next(k for k, t in enumerate(P) if k > i_abs and t.startswith('Keywords'))])
nres, nabs = len(res.split()), len(abs_.split())
check('R1  resumen de 250 a 300 palabras', 250 <= nres <= 300, 'tiene %d' % nres)
check('R2  abstract espejo (230 a 320 palabras)', 230 <= nabs <= 320, 'tiene %d' % nabs)
for rot, cita, ref in [('Wilson', 'Wilson, 1927', 'Wilson, E. B. (1927)'),
                       ('McNemar', 'McNemar, 1947', 'McNemar, Q. (1947)'),
                       ('Holm', 'Holm (1979)', 'Holm, S. (1979)'),
                       ('Mann-Whitney', 'Mann y Whitney, 1947', 'Mann, H. B., & Whitney, D. R. (1947)'),
                       ('Welch', 'Welch (1947)', 'Welch, B. L. (1947)'),
                       ('Fieller', 'Fieller (1954)', 'Fieller, E. C. (1954)'),
                       ('Cohen', 'Cohen (1960)', 'Cohen, J. (1960)'),
                       ('CLINC150', 'Larson et al., 2019', 'Larson, S.'),
                       ('Banking77', 'Casanueva et al., 2020', 'Casanueva, I.'),
                       ('SNIPS', 'Coucke et al., 2018', 'Coucke, A.')]:
    check('R3  %-13s citado y en la lista' % rot, cita in CUERPO and any(r.startswith(ref) for r in REFS))
narr = [m for p in P[:i_ref] for m in re.findall(r'[A-ZÁÉÍÓÚ][\w-]+ & [A-ZÁÉÍÓÚ][\w-]+ \(\d{4}\)', p)]
check('R4  APA: et al. desde la primera cita y «y» en citas narrativas',
      not narr and 'Alderete, Jones y Motta' not in CUERPO and 'Fondevila-Gascón, Huamanchumo' not in CUERPO
      and 'Pachas-Santos, Calderón-Vilca' not in CUERPO and 'Bravo Maruri, Ramírez Reina' not in CUERPO, str(narr))
b541 = bloque('5.4.1', 'CAPÍTULO 6')
etiquetas = re.findall(r'^\(([a-z])(?:-\w+)?\)', b541, re.M)
check('R5  §5.4.1 sin bis ni ter y en orden',
      '-bis)' not in TODO and '-ter)' not in TODO
      and etiquetas == [chr(ord('a') + k) for k in range(len(etiquetas))] and len(etiquetas) >= 9, str(etiquetas))
b357 = bloque('3.5.7', '3.6 Amenazas')
b525 = bloque('5.2.5', '5.3 ')
check('R6  E6: piloto y alcance de la evaluación declarados', 'prueba piloto' in b357 and 'GENERAL' in b357)
check('R7  E6: horarios en las reglas y desobediencia a «consultar con el equipo»',
      'horarios de atención' in b525 and 'consultar con el equipo' in b525)
check('R8  §3.6 ampliada (temperatura, alias, operador único)',
      all(x in b36 for x in ('temperatura', 'alias', 'único operador')))
CAP7 = bloque('CAPÍTULO 7', 'CAPÍTULO 8')
check('R9  líneas futuras nuevas',
      all(x in CAP7 for x in ('Tiendanube', 'inyección de instrucciones', 'atención manual por chat', 'Piloto en una PyME')))
check('R10 la línea de ajuste fino revisada',
      'Fine-tuning del modelo de IA con datos reales de conversaciones de clientes para mejorar la precisión' not in TODO)
check('R11 la pérdida silenciosa de interacciones en limitaciones y recomendaciones',
      'fuera del vocabulario' in b541 and 'fuera del vocabulario' in CAP7)
# el estilo Source Code corta palabras al azar (wordWrap=0); los prompts del Anexo H son prosa y deben cortar por palabra
prompts_h = [p for p in d.paragraphs if p.style.name == 'Source Code' and p.text.startswith(('=# --- START', '# --- START'))]
check('R13 los prompts del Anexo H cortan línea por palabra',
      len(prompts_h) == 2 and all(re.search(r'<w:wordWrap(?: w:val="(?:1|true|on)")?/>', p._p.pPr.xml)
                                  for p in prompts_h))   # Word guarda el verdadero como <w:wordWrap/>
lst = [f for t in d.tables for f in filas(t) if f and f[0] == 'Tabla 5.10']
cap = next((p for p in P if p.startswith('Tabla 5.10:')), '')
check('R12 el listado de tablas coincide con el epígrafe de la Tabla 5.10',
      bool(lst) and cap[len('Tabla 5.10: '):].rstrip('.') == lst[0][1], '%s | %s' % (cap[:80], lst[:1]))

print()
print('=' * 78)
if fallos:
    print(' RESULTADO: %d/%d — %d FALLAS' % (n_ok, n_ok + len(fallos), len(fallos)))
    print('=' * 78)
    sys.exit(1)
print(' DICTAMEN DEL 14/09: %d/%d' % (n_ok, n_ok))
print('=' * 78)
